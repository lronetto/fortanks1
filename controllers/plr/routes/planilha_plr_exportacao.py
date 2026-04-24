"""Exportações Excel do relatório planilha PLR."""

from io import BytesIO

from flask import flash, redirect, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .. import plr_bp
from ..services.planilha_plr_excel import (
    _excel_aba_resumo,
    _excel_aba_resumo_formulas,
    _excel_abas_meses,
    _excel_abas_meses_formulas,
    _excel_estilos,
)
from ..services.planilha_plr_calculo import MESES_ABREV
from ..services.planilha_plr_template import gerar_planilha_plr_xlsx_template_io
from ..services.relatorio_planilha_payload import carregar_payload_planilha, parse_filtros_planilha


@plr_bp.route('/relatorio-planilha/excel')
def relatorio_planilha_excel():
    """Gera arquivo Excel com aba resumo e uma aba por mês, usando os mesmos filtros do relatório."""
    err, filtros = parse_filtros_planilha(request.values)
    if err:
        flash(err, 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        payload = carregar_payload_planilha(filtros)
    except Exception as e:
        flash(f'Erro ao gerar Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    resultado = payload['resultado']
    meses_colunas = payload['meses_colunas']
    data_fechamento = payload['data_fechamento']
    data_inicio = payload['data_inicio']
    cpf_para_chapa = payload['cpf_para_chapa']
    criterios_por_colab_mes = payload['criterios_por_colab_mes']
    fi = filtros

    wb = Workbook()
    estilos = _excel_estilos()
    _excel_aba_resumo(wb, resultado, meses_colunas, cpf_para_chapa, estilos)
    _excel_abas_meses(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Planilha_{fi["ano_inicio"]}-{fi["mes_inicio"]}_a_{fi["ano_fim"]}-{fi["mes_fim"]}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@plr_bp.route('/relatorio-planilha/excel-formulas')
def relatorio_planilha_excel_formulas():
    """
    Gera o mesmo Excel do relatório planilha, porém com SOMA, P, VPO e Valor total
    calculados por fórmulas do Excel (e totalizador em fórmula).
    """
    err, filtros = parse_filtros_planilha(request.values)
    if err:
        flash(err, 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        payload = carregar_payload_planilha(filtros)
    except Exception as e:
        flash(f'Erro ao gerar Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    resultado = payload['resultado']
    meses_colunas = payload['meses_colunas']
    cpf_para_chapa = payload['cpf_para_chapa']
    criterios_por_colab_mes = payload['criterios_por_colab_mes']
    fi = filtros

    wb = Workbook()
    wb.remove(wb.worksheets[0])
    estilos = _excel_estilos()
    _excel_abas_meses_formulas(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos)
    _excel_aba_resumo_formulas(wb, resultado, meses_colunas, cpf_para_chapa, estilos)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Planilha_Formulas_{fi["ano_inicio"]}-{fi["mes_inicio"]}_a_{fi["ano_fim"]}-{fi["mes_fim"]}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@plr_bp.route('/relatorio-planilha/excel-planilha-plr')
def relatorio_planilha_excel_planilha_plr():
    """
    Exporta Excel usando o template PLANILHA_PLR.xlsx (abas por mês + PAGAMENTO MOD).
    Mesmos filtros do relatório planilha.
    """
    err, filtros = parse_filtros_planilha(request.values)
    if err:
        flash(err, 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        payload = carregar_payload_planilha(filtros)
    except Exception as e:
        flash(f'Erro ao gerar Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    resultado = payload['resultado']
    meses_colunas = payload['meses_colunas']
    data_fechamento = payload['data_fechamento']
    data_inicio = payload['data_inicio']
    cpf_para_chapa = payload['cpf_para_chapa']
    criterios_por_colab_mes = payload['criterios_por_colab_mes']
    fi = filtros
    salario_grupo = bool(fi.get('salario_por_grupo'))

    try:
        output = gerar_planilha_plr_xlsx_template_io(
            resultado,
            meses_colunas,
            data_inicio,
            data_fechamento,
            cpf_para_chapa,
            criterios_por_colab_mes,
            salario_por_grupo=salario_grupo,
        )
    except FileNotFoundError as e:
        flash(str(e), 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    except Exception as e:
        flash(f'Erro ao montar planilha pelo template: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    filename = f'PLR_Planilha_MOD_{fi["ano_inicio"]}-{fi["mes_inicio"]}_a_{fi["ano_fim"]}-{fi["mes_fim"]}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@plr_bp.route('/relatorio-planilha/assiduidade/template')
def relatorio_planilha_assiduidade_template():
    """Gera Excel template para importação de assiduidade: linhas = colaboradores (CPF, Nome), colunas = meses (faltas)."""
    err, filtros = parse_filtros_planilha(request.values)
    if err:
        flash(err, 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        payload = carregar_payload_planilha(filtros)
    except Exception as e:
        flash(f'Erro ao gerar template: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    resultado = payload['resultado']
    meses_colunas = payload['meses_colunas']
    fi = filtros

    wb = Workbook()
    ws = wb.active
    ws.title = 'Assiduidade'
    header = ['CPF', 'Nome']
    for (mes, ano) in meses_colunas:
        idx = int(mes) - 1
        mes_nome = MESES_ABREV[idx] if 0 <= idx < len(MESES_ABREV) else f'{mes:02d}'
        header.append(f'{mes_nome} {ano}')
    ws.append(header)
    for r in resultado:
        colab = r.get('colaborador')
        if not colab:
            continue
        cpf = colab.cpf or ''
        nome = (colab.nome or '').upper()
        row = [cpf, nome]
        for _ in meses_colunas:
            row.append(0)
        ws.append(row)
    estilos = _excel_estilos()
    for cell in ws[1]:
        cell.font = estilos['header_font']
        cell.fill = estilos['header_fill']
        cell.alignment = estilos['center_align']
    for col_idx, _ in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14 if col_idx <= 2 else 7.3
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Assiduidade_Template_{fi["ano_inicio"]}-{fi["mes_inicio"]}_a_{fi["ano_fim"]}-{fi["mes_fim"]}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )
