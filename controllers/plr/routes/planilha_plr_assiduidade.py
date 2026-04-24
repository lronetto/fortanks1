"""Preview e importação de assiduidade (planilha PLR)."""

from flask import flash, jsonify, redirect, request, url_for

from models.database import db
from models.plr import PlrAssiduidade

from .. import plr_bp
from ..services.planilha_plr_assiduidade import mapa_cpf_colaborador, normalizar_cpf, parse_meses_do_header
from ..services.planilha_plr_calculo import MESES_ABREV


@plr_bp.route('/relatorio-planilha/assiduidade/preview-meses', methods=['POST'])
def relatorio_planilha_assiduidade_preview_meses():
    """Retorna JSON com os meses/anos detectados no cabeçalho do arquivo Excel enviado."""
    arquivo = request.files.get('arquivo') or request.files.get('file')
    if not arquivo or not arquivo.filename:
        return jsonify({'ok': False, 'error': 'Nenhum arquivo enviado.'}), 400
    try:
        from openpyxl import load_workbook

        wb = load_workbook(arquivo, read_only=True, data_only=True)
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'Erro ao ler o Excel: {e}'}), 400
    if not header:
        return jsonify({'ok': False, 'error': 'Planilha sem cabeçalho.'}), 400

    meses = parse_meses_do_header(list(header))
    if not meses:
        return jsonify({'ok': False, 'error': 'Nenhum mês/ano identificado nos cabeçalhos.'}), 400

    labels = [f'{MESES_ABREV[m - 1]}/{a}' for m, a in meses]
    return jsonify({'ok': True, 'meses': labels})


@plr_bp.route('/relatorio-planilha/assiduidade/import', methods=['POST'])
def relatorio_planilha_assiduidade_import():
    """Importa planilha Excel de assiduidade (faltas por colaborador por mês) e grava em PlrAssiduidade."""
    if 'arquivo' not in request.files and 'file' not in request.files:
        flash('Nenhum arquivo enviado.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    arquivo = request.files.get('arquivo') or request.files.get('file')
    if not arquivo or not arquivo.filename or not arquivo.filename.lower().endswith(('.xlsx', '.xls')):
        flash('Envie um arquivo Excel (.xlsx).', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    try:
        from openpyxl import load_workbook

        wb = load_workbook(arquivo, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        flash(f'Erro ao ler o Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    if len(rows) < 2:
        flash('Planilha deve ter cabeçalho e ao menos uma linha de dados.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    header = rows[0]
    meses_colunas = parse_meses_do_header(header)

    if not meses_colunas:
        flash(
            'Não foi possível identificar os meses/ano nos cabeçalhos da planilha. '
            'Use o formato do template (ex: JAN 2025, FEV 2025).',
            'danger',
        )
        return redirect(url_for('plr.relatorio_planilha'))

    meses_label = ', '.join(f'{MESES_ABREV[m - 1]}/{a}' for m, a in meses_colunas)

    colunas_meses = len(meses_colunas)
    inseridos = 0
    atualizados = 0
    erros = []
    cpf_para_colab = mapa_cpf_colaborador()

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or len(row) < 2:
            continue
        cpf_cel = row[0]
        cpf_limpo = normalizar_cpf(cpf_cel)
        if not cpf_limpo or len(cpf_limpo) < 11:
            continue
        colab = cpf_para_colab.get(cpf_limpo)
        if not colab:
            erros.append(f'Linha {row_idx}: CPF não encontrado ({cpf_cel})')
            continue
        for col_idx in range(colunas_meses):
            if col_idx + 2 >= len(row):
                break
            mes, ano = meses_colunas[col_idx]
            val = row[col_idx + 2]
            try:
                faltas = int(float(val)) if val is not None else 0
            except (TypeError, ValueError):
                faltas = 0
            faltas = max(0, faltas)
            rec = PlrAssiduidade.query.filter_by(
                colaborador_id=colab.id, mes=mes, ano=ano
            ).first()
            if rec:
                rec.faltas = faltas
                atualizados += 1
            else:
                db.session.add(PlrAssiduidade(colaborador_id=colab.id, mes=mes, ano=ano, faltas=faltas))
                inseridos += 1
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao salvar: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    msg_periodo = f' Período detectado: {meses_label}.'
    if erros:
        flash(
            f'Importado: {inseridos} novos, {atualizados} atualizados.{msg_periodo} Avisos: '
            + '; '.join(erros[:5]),
            'warning',
        )
    else:
        flash(f'Assiduidade importada: {inseridos} novos, {atualizados} atualizados.{msg_periodo}', 'success')
    next_url = (request.form.get('next') or '').strip()
    if next_url and next_url.startswith('/'):
        return redirect(next_url)
    return redirect(url_for('plr.relatorio_planilha'))
