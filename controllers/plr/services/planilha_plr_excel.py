"""
Geração de workbooks Excel do relatório planilha PLR (abas resumo / meses / fórmulas).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from models.database import db
from models.plr import PLRColaborador, PlrAssiduidade
from models.colaborador import Colaborador

from models.plr.utils import assiduidade_pct_por_faltas

from ..constantes import CRITERIOS_PADRAO_PLR
from .avaliacoes import valor_por_tipo as _valor_por_tipo
from .planilha_plr_calculo import (
    MESES_ABREV,
    MIN_MESES_AVALIACAO_RESUMO,
    _titulo_aba_mes,
    _criterios_headers_com_peso,
)


def _valor_criterio_para_excel_pct(val):
    """
    Converte valor do critério para número 0-1 para Excel (formato %).
    Se valor está em escala 0-10 (nota), converte para 0-100; se já é 0-100, mantém.
    Retorna None se val for None.
    """
    if val is None:
        return None
    try:
        v = float(val)
    except (TypeError, ValueError):
        return None
    if v <= 10:
        v = v * 10.0
    return min(100.0, max(0.0, v)) / 100.0


def _pesos_por_tipo():
    """Retorna dict tipo -> peso (0-1) a partir de CRITERIOS_PADRAO_PLR."""
    return {c['tipo']: (c.get('peso') or 0) for c in CRITERIOS_PADRAO_PLR}


def _valor_criterio_ponderado_excel(val, peso):
    """
    Retorna o valor do critério já ponderado pelo peso (valor * peso) para Excel em formato %.
    Ex.: valor 85% (0.85) com peso 30% (0.30) -> 0.255 (exibe 25.5%).
    """
    if val is None or peso is None:
        return None
    pct = _valor_criterio_para_excel_pct(val)
    if pct is None:
        return None
    try:
        p = float(peso)
    except (TypeError, ValueError):
        return None
    return pct * p


def _excel_estilos():
    """Estilos reutilizáveis para cabeçalho e células."""
    return {
        'header_font': Font(bold=True, color="FFFFFF"),
        'header_fill': PatternFill("solid", fgColor="4F81BD"),
        'center_align': Alignment(horizontal="center", vertical="center", wrap_text=True),
        'right_align': Alignment(horizontal="right", vertical="center"),
        'thin_border': Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        ),
    }


def _excel_aba_resumo(wb, resultado, meses_colunas, cpf_para_chapa, estilos):
    """Preenche e formata a aba Resumo."""
    ws = wb.active
    ws.title = 'Resumo'
    header = ['CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão/Fech.', 'Tempo casa (meses)', 'Função']
    for (mes, ano) in meses_colunas:
        header.append(f'{_titulo_aba_mes(mes, ano)}')
    header.extend(['SOMA (%)', 'P (%)', 'Salário base PLR', 'VPO', 'Valor total'])
    ws.append(header)

    total_geral = 0
    for r in resultado:
        pcts_meses = r.get('pcts_meses') or []
        meses_com_avaliacao = [pct for pct in pcts_meses if pct is not None]
        if len(meses_com_avaliacao) < MIN_MESES_AVALIACAO_RESUMO:
            continue
        total_geral += (r.get('valor_total') or 0)
        colab = r['colaborador']
        data_fech = r.get('data_fechamento')
        cpf = colab.cpf or '-'
        chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
        admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else '-'
        if colab.data_demissao and data_fech and colab.data_demissao <= data_fech:
            demissao_fech = colab.data_demissao.strftime('%d/%m/%Y')
        else:
            demissao_fech = data_fech.strftime('%d/%m/%Y') if data_fech else '-'
        funcao = colab.cargo.nome if colab.cargo else '-'
        valores_mensais = [
            (pcts_meses[i] / 100.0) if i < len(pcts_meses) and pcts_meses[i] is not None else None
            for i in range(len(meses_colunas))
        ]
        linha = (
            [cpf, chapa, (colab.nome or '').upper(), admissao, demissao_fech, r['tempo_casa_meses'], funcao]
            + valores_mensais
            + [
                (r['soma'] / 100.0) if r['soma'] is not None else None,
                (r['p'] / 100.0) if r['p'] is not None else None,
                r['salario_base_plr'], r['vpo'], r['valor_total'],
            ]
        )
        ws.append(linha)

    # Linha de total: mescla células e exibe totalizador do valor total
    num_cols = len(header)
    total_row = ws.max_row + 1
    row_data = ['TOTAL'] + [None] * (num_cols - 2) + [total_geral]
    ws.append(row_data)
    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=num_cols - 1,
    )
    cell_total_label = ws.cell(row=total_row, column=1)
    cell_total_label.value = 'TOTAL'
    cell_total_label.font = Font(bold=True)
    cell_total_label.alignment = estilos['right_align']
    cell_total_label.border = estilos['thin_border']
    cell_total_val = ws.cell(row=total_row, column=num_cols)
    cell_total_val.number_format = r'R$ #,##0.00'
    cell_total_val.font = Font(bold=True)
    cell_total_val.alignment = estilos['right_align']
    cell_total_val.border = estilos['thin_border']
    for col in range(2, num_cols):
        c = ws.cell(row=total_row, column=col)
        c.border = estilos['thin_border']

    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = estilos['header_font']
            cell.fill = estilos['header_fill']
            cell.alignment = estilos['center_align']
            cell.border = estilos['thin_border']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = estilos['thin_border']

    num_meses = len(meses_colunas)
    col_mes_inicio = 8
    col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 12, 'G': 29}
    for idx_col in range(col_mes_inicio, col_mes_inicio + num_meses):
        col_letter = ws.cell(row=1, column=idx_col).column_letter
        col_widths[col_letter] = 7.3
    for idx_col in range(col_mes_inicio + num_meses, ws.max_column + 1):
        col_letter = ws.cell(row=1, column=idx_col).column_letter
        col_widths[col_letter] = 14
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    right = estilos['right_align']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = right
    if num_meses:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_mes_inicio, max_col=col_mes_inicio + num_meses - 1):
            for cell in row:
                cell.number_format = "0.00%"
                cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
        for cell in col:
            cell.number_format = "0"
            cell.alignment = right
    soma_col = col_mes_inicio + num_meses + 1
    p_col = soma_col + 1
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=soma_col, max_col=p_col):
        for cell in col:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col_idx in (p_col + 1, p_col + 2, p_col + 3):
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in col:
                cell.number_format = r'R$ #,##0.00'
                cell.alignment = right


def _excel_aba_resumo_formulas(wb, resultado, meses_colunas, cpf_para_chapa, estilos):
    """
    Preenche a aba Resumo inserida no início do workbook. O % de cada mês vem da aba
    do mês correspondente (fórmula INDEX/MATCH por CPF). SOMA, P, VPO e Valor total em fórmulas.
    """
    ws = wb.create_sheet('Resumo', 0)
    num_meses = len(meses_colunas)
    col_mes_inicio = 8
    col_mes_fim = 7 + num_meses
    col_soma = 8 + num_meses
    col_p = 9 + num_meses
    col_salario = 10 + num_meses
    col_vpo = 11 + num_meses
    col_valor_total = 12 + num_meses
    col_mes_pct_na_aba_mes = 11

    header = ['CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão/Fech.', 'Tempo casa (meses)', 'Função']
    for (mes, ano) in meses_colunas:
        header.append(_titulo_aba_mes(mes, ano))
    header.extend(['SOMA (%)', 'P (%)', 'Salário base PLR', 'VPO', 'Valor total'])
    ws.append(header)

    range_mes_ini = get_column_letter(col_mes_inicio)
    range_mes_fim = get_column_letter(col_mes_fim)
    letter_salario = get_column_letter(col_salario)
    letter_p = get_column_letter(col_p)
    letter_vpo = get_column_letter(col_vpo)
    letter_valor_total = get_column_letter(col_valor_total)
    col_mes_letter_aba = get_column_letter(col_mes_pct_na_aba_mes)

    data_row = 1
    for r in resultado:
        if False: 
            pcts_meses = r.get('pcts_meses') or []
            meses_com_avaliacao = [pct for pct in pcts_meses if pct is not None]
            if len(meses_com_avaliacao) < MIN_MESES_AVALIACAO_RESUMO:
                continue
        data_row += 1
        row_idx = data_row
        colab = r['colaborador']
        data_fech = r.get('data_fechamento')
        cpf = colab.cpf or '-'
        chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
        admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else '-'
        if colab.data_demissao and data_fech and colab.data_demissao <= data_fech:
            demissao_fech = colab.data_demissao.strftime('%d/%m/%Y')
        else:
            demissao_fech = data_fech.strftime('%d/%m/%Y') if data_fech else '-'
        funcao = colab.cargo.nome if colab.cargo else '-'
        for col, val in enumerate([cpf, chapa, (colab.nome or '').upper(), admissao, demissao_fech, r['tempo_casa_meses'], funcao], start=1):
            ws.cell(row=row_idx, column=col, value=val)
        for i, (mes, ano) in enumerate(meses_colunas):
            sheet_mes = f"'{_titulo_aba_mes(mes, ano)}'"
            col_resumo = col_mes_inicio + i
            ws.cell(
                row=row_idx,
                column=col_resumo,
                value=f"=IFERROR(INDEX({sheet_mes}!${col_mes_letter_aba}:${col_mes_letter_aba},MATCH($A{row_idx},{sheet_mes}!$A:$A,0)),\"\")"
            )
        ws.cell(row=row_idx, column=col_salario, value=r.get('salario_base_plr'))
        ws.cell(row=row_idx, column=col_soma, value=f'=SUM({range_mes_ini}{row_idx}:{range_mes_fim}{row_idx})')
        ws.cell(row=row_idx, column=col_p, value=f'=SUM({range_mes_ini}{row_idx}:{range_mes_fim}{row_idx})/{num_meses}')
        ws.cell(row=row_idx, column=col_vpo, value=f'=({letter_salario}{row_idx}/12)*6')
        ws.cell(row=row_idx, column=col_valor_total, value=f'={letter_vpo}{row_idx}*{letter_p}{row_idx}')

    last_data_row = ws.max_row
    total_row = last_data_row + 1
    num_cols = col_valor_total
    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=num_cols - 1,
    )
    cell_total_label = ws.cell(row=total_row, column=1)
    cell_total_label.value = 'TOTAL'
    cell_total_label.font = Font(bold=True)
    cell_total_label.alignment = estilos['right_align']
    cell_total_label.border = estilos['thin_border']
    cell_total_val = ws.cell(row=total_row, column=num_cols)
    cell_total_val.value = f'=SUM({letter_valor_total}2:{letter_valor_total}{last_data_row})'
    cell_total_val.number_format = r'R$ #,##0.00'
    cell_total_val.font = Font(bold=True)
    cell_total_val.alignment = estilos['right_align']
    cell_total_val.border = estilos['thin_border']
    for col in range(2, num_cols):
        ws.cell(row=total_row, column=col).border = estilos['thin_border']

    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = estilos['header_font']
        cell.fill = estilos['header_fill']
        cell.alignment = estilos['center_align']
        cell.border = estilos['thin_border']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = estilos['thin_border']

    num_meses_resumo = len(meses_colunas)
    col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 12, 'G': 29}
    for idx_col in range(col_mes_inicio, col_mes_inicio + num_meses_resumo):
        col_widths[get_column_letter(idx_col)] = 7.3
    for idx_col in range(col_mes_inicio + num_meses_resumo, ws.max_column + 1):
        col_widths[get_column_letter(idx_col)] = 14
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    right = estilos['right_align']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = right
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_mes_inicio, max_col=col_mes_fim):
        for cell in row:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
        for cell in col:
            cell.number_format = "0"
            cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_soma, max_col=col_p):
        for cell in col:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col_idx in (col_salario, col_vpo, col_valor_total):
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in col:
                cell.number_format = r'R$ #,##0.00'
                cell.alignment = right


def _excel_abas_meses_formulas(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos):
    """
    Cria uma aba por mês com fórmulas: % Mês = soma ponderada dos critérios,
    Qtd Meses e Acum (%) calculados por fórmulas com referências cruzadas entre abas.
    O Resumo usa o % Mês destas abas (fórmula INDEX/MATCH por CPF).
    """
    header_font = estilos['header_font']
    header_fill = estilos['header_fill']
    center_align = estilos['center_align']
    right_align = estilos['right_align']
    thin_border = estilos['thin_border']
    col_pct_mes = 11
    col_qtd = 12
    col_acum = 13
    letter_pct_mes = get_column_letter(col_pct_mes)

    headers_criterios = [label for _, label in _criterios_headers_com_peso()]
    pesos = _pesos_por_tipo()

    titulos_abas = [_titulo_aba_mes(m, a) for m, a in meses_colunas]

    for idx, (mes, ano) in enumerate(meses_colunas):
        titulo = titulos_abas[idx]
        ws = wb.create_sheet(title=titulo)
        header_mes = [
            'CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão', 'Função',
        ] + headers_criterios + [
            '% Mês', 'Qtd Meses', 'Acum (%)',
        ]
        ws.append(header_mes)

        data_row = 1
        for r in resultado:
            pct_mes_list = r.get('pcts_meses') or []
            pct_mes = pct_mes_list[idx] if idx < len(pct_mes_list) else None
            if pct_mes is None:
                continue
            data_row += 1
            row_idx = data_row
            colab = r['colaborador']
            cpf = colab.cpf or '-'
            chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
            admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else ''
            demissao = colab.data_demissao.strftime('%d/%m/%Y') if colab.data_demissao else ''
            funcao = colab.cargo.nome if colab.cargo else '-'

            crit_colab = criterios_por_colab_mes.get(colab.id, {})
            crit_mes = crit_colab.get((mes, ano), {}) if crit_colab else {}
            assid = _valor_criterio_ponderado_excel(crit_mes.get('Assiduidade'), pesos.get('Assiduidade'))
            zero_ac = _valor_criterio_ponderado_excel(crit_mes.get('Zero Acidente'), pesos.get('Zero Acidente'))
            segur = _valor_criterio_ponderado_excel(crit_mes.get('Segurança, Limpeza, Organização'), pesos.get('Segurança, Limpeza, Organização'))
            prazo = _valor_criterio_ponderado_excel(crit_mes.get('Prazo'), pesos.get('Prazo'))

            for col, val in enumerate([cpf, chapa, (colab.nome or '').upper(), admissao, demissao, funcao, assid, zero_ac, segur, prazo], start=1):
                ws.cell(row=row_idx, column=col, value=val)

            ws.cell(row=row_idx, column=col_pct_mes,
                    value=f'=G{row_idx}+H{row_idx}+I{row_idx}+J{row_idx}')

            qtd_parts = []
            sum_parts = []
            for j in range(idx + 1):
                if j == idx:
                    qtd_parts.append(f'IF({letter_pct_mes}{row_idx}>0,1,0)')
                    sum_parts.append(f'{letter_pct_mes}{row_idx}')
                else:
                    tab_ref = f"'{titulos_abas[j]}'"
                    qtd_parts.append(
                        f'IF(ISNUMBER(MATCH($A{row_idx},{tab_ref}!$A:$A,0)),1,0)'
                    )
                    sum_parts.append(
                        f'IFERROR(INDEX({tab_ref}!${letter_pct_mes}:${letter_pct_mes},MATCH($A{row_idx},{tab_ref}!$A:$A,0)),0)'
                    )

            ws.cell(row=row_idx, column=col_qtd,
                    value=f'={"+".join(qtd_parts)}')
            num_meses_periodo = idx + 1
            ws.cell(row=row_idx, column=col_acum,
                    value=f'=({"+".join(sum_parts)})/{num_meses_periodo}')

        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border

        col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 29, 'G': 14, 'H': 14, 'I': 22, 'J': 12, 'K': 10, 'L': 10, 'M': 10}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        for col_idx in range(4, 6):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "dd/mm/yyyy"
                    cell.alignment = right_align
        for col_idx in range(7, 11):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.0%"
                    cell.alignment = right_align
        for col_idx in (col_pct_mes, col_acum):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.00%"
                    cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_qtd, max_col=col_qtd):
            for cell in col:
                cell.number_format = "0"
                cell.alignment = right_align


def _excel_abas_meses(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos):
    """Cria uma aba por mês com colaboradores que têm avaliação no mês."""
    header_font = estilos['header_font']
    header_fill = estilos['header_fill']
    center_align = estilos['center_align']
    right_align = estilos['right_align']
    thin_border = estilos['thin_border']
    headers_criterios = [label for _, label in _criterios_headers_com_peso()]
    pesos = _pesos_por_tipo()

    for idx, (mes, ano) in enumerate(meses_colunas):
        titulo = _titulo_aba_mes(mes, ano)
        ws = wb.create_sheet(title=titulo)
        header_mes = [
            'CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão', 'Função',
        ] + headers_criterios + [
            '% Mês', 'Qtd Meses', 'Acum (%)',
        ]
        ws.append(header_mes)

        for r in resultado:
            colab = r['colaborador']
            pct_mes_list = r.get('pcts_meses') or []
            pct_mes = pct_mes_list[idx] if idx < len(pct_mes_list) else None
            if pct_mes is None:
                continue
            valid_prev = [p for j, p in enumerate(pct_mes_list) if j <= idx and p is not None]
            qtd_meses = len(valid_prev)
            num_meses_periodo = idx + 1
            acum_pct = (sum(valid_prev) / num_meses_periodo) if num_meses_periodo else None

            cpf = colab.cpf or '-'
            chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
            admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else ''
            demissao = colab.data_demissao.strftime('%d/%m/%Y') if colab.data_demissao else ''
            funcao = colab.cargo.nome if colab.cargo else '-'
            crit_colab = criterios_por_colab_mes.get(colab.id, {})
            crit_mes = crit_colab.get((mes, ano), {}) if crit_colab else {}
            assid = _valor_criterio_ponderado_excel(crit_mes.get('Assiduidade'), pesos.get('Assiduidade'))
            zero_ac = _valor_criterio_ponderado_excel(crit_mes.get('Zero Acidente'), pesos.get('Zero Acidente'))
            segur = _valor_criterio_ponderado_excel(crit_mes.get('Segurança, Limpeza, Organização'), pesos.get('Segurança, Limpeza, Organização'))
            prazo = _valor_criterio_ponderado_excel(crit_mes.get('Prazo'), pesos.get('Prazo'))

            ws.append([
                cpf, chapa, (colab.nome or '').upper(), admissao, demissao, funcao,
                assid, zero_ac, segur, prazo,
                (pct_mes / 100.0),
                qtd_meses,
                (acum_pct / 100.0) if acum_pct is not None else None,
            ])

        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border

        col_widths_mes = {
            'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 29,
            'G': 14, 'H': 14, 'I': 22, 'J': 12, 'K': 10, 'L': 10, 'M': 10,
        }
        for col, width in col_widths_mes.items():
            ws.column_dimensions[col].width = width

        for col_idx in range(4, 6):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "dd/mm/yyyy"
                    cell.alignment = right_align
        for col_idx in range(7, 11):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.0%"
                    cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=11, max_col=11):
            for cell in col:
                cell.number_format = "0.00%"
                cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=12, max_col=12):
            for cell in col:
                cell.number_format = "0"
                cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=13, max_col=13):
            for cell in col:
                cell.number_format = "0.00%"
                cell.alignment = right_align
