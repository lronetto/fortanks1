"""
Relatório Planilha PLR: cálculo por período, dados JSON (DataTables), export Excel e página.
"""
from datetime import date, timedelta
from io import BytesIO

from flask import request, render_template, redirect, url_for, flash, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from models.database import db
from models.plr import PLRColaborador, EfetivoPLR, PlrAssiduidade
from models.colaborador import Colaborador
from models.departamento import Departamento

from utils.plr_calculo import assiduidade_pct_por_faltas

from . import plr_bp
from .avaliacao import _equipes_distintas_plr, _valor_por_tipo
from .constantes import CRITERIOS_PADRAO_PLR


MESES_ABREV = ['JAN', 'FEV', 'MAR', 'ABR', 'MAI', 'JUN', 'JUL', 'AGO', 'SET', 'OUT', 'NOV', 'DEZ']


def _titulo_aba_mes(mes, ano):
    """Retorna o título da aba do mês no formato 'AGO 2025', 'SET 2025', etc."""
    idx = int(mes) - 1
    mes_nome = MESES_ABREV[idx] if 0 <= idx < len(MESES_ABREV) else f'{mes:02d}'
    return f'{mes_nome} {ano}'


MIN_MESES_AVALIACAO_RESUMO = 3

# Títulos das colunas de avaliação com peso em % (cache)
CRITERIOS_HEADERS_COM_PESO = None


def _criterios_headers_com_peso():
    """Retorna lista [('Assiduidade', 'ASSIDUIDADE (30%)'), ...] com pesos do CRITERIOS_PADRAO_PLR (títulos em maiúsculas)."""
    global CRITERIOS_HEADERS_COM_PESO
    if CRITERIOS_HEADERS_COM_PESO is not None:
        return CRITERIOS_HEADERS_COM_PESO
    out = []
    for c in CRITERIOS_PADRAO_PLR:
        tipo = c.get('tipo', '')
        peso = c.get('peso')
        pct = int(round((peso * 100))) if peso is not None else ''
        label = f'{tipo.upper()} ({pct}%)' if pct != '' else tipo.upper()
        out.append((tipo, label))
    CRITERIOS_HEADERS_COM_PESO = out
    return out


def _valor_criterio_para_excel_pct(val):
    """
    Converte valor do critério para número 0-1 para Excel (formato %).
    Se valor está em escala 0-10 (nota), converte para 0-100; se já é 0-100, mantém.
    Retorna None se val for None.
    """
    if val is None:
        return None
    try:
        v = float(val)
    except (TypeError, ValueError):
        return None
    if v <= 10:
        v = v * 10.0
    return min(100.0, max(0.0, v)) / 100.0


def _pesos_por_tipo():
    """Retorna dict tipo -> peso (0-1) a partir de CRITERIOS_PADRAO_PLR."""
    return {c['tipo']: (c.get('peso') or 0) for c in CRITERIOS_PADRAO_PLR}


def _valor_criterio_ponderado_excel(val, peso):
    """
    Retorna o valor do critério já ponderado pelo peso (valor * peso) para Excel em formato %.
    Ex.: valor 85% (0.85) com peso 30% (0.30) -> 0.255 (exibe 25.5%).
    """
    if val is None or peso is None:
        return None
    pct = _valor_criterio_para_excel_pct(val)
    if pct is None:
        return None
    try:
        p = float(peso)
    except (TypeError, ValueError):
        return None
    return pct * p


def relatorio_planilha_calcular(ano_inicio, mes_inicio, ano_fim, mes_fim, modelo_plr_id,
                                equipe_filtro=None, departamentos_ids=None):
    """
    Retorna (resultado, meses_colunas, data_fechamento) para o relatório planilha.
    Suporta períodos que abrangem anos diferentes.
    departamentos_ids: lista de IDs de departamento para filtrar (None ou vazia = todos).
    """
    from utils.plr_calculo import tempo_de_casa_meses, salario_base_plr, nota_media_com_assiduidade

    mes_i = int(mes_inicio)
    mes_f = int(mes_fim)
    ano_i = int(ano_inicio)
    ano_f = int(ano_fim)
    data_inicio = date(ano_i, mes_i, 1)
    data_fechamento = date(ano_f, mes_f, 1) + timedelta(days=32)
    data_fechamento = data_fechamento.replace(day=1) - timedelta(days=1)

    meses_colunas = []
    cur_ano, cur_mes = ano_i, mes_i
    while (cur_ano < ano_f) or (cur_ano == ano_f and cur_mes <= mes_f):
        meses_colunas.append((cur_mes, cur_ano))
        cur_mes += 1
        if cur_mes > 12:
            cur_mes = 1
            cur_ano += 1

    q_av = PLRColaborador.query.filter(
        PLRColaborador.data >= data_inicio,
        PLRColaborador.data <= data_fechamento,
    )
    if modelo_plr_id:
        q_av = q_av.filter(PLRColaborador.PlrModelo_id == int(modelo_plr_id))
    avaliacoes_periodo = q_av.all()

    if equipe_filtro:
        avaliacoes_periodo = [
            av for av in avaliacoes_periodo
            if av.equipe_alocada and isinstance(av.equipe_alocada, list)
            and equipe_filtro in [str(e).strip() for e in av.equipe_alocada if e]
        ]

    ids_colab = list({av.colaborador_id for av in avaliacoes_periodo})
    if not ids_colab:
        colaboradores = (
            Colaborador.query.filter(Colaborador.data_admissao <= data_fechamento)
            .filter(
                (Colaborador.data_demissao.is_(None)) | (Colaborador.data_demissao >= data_inicio)
            )
            .order_by(Colaborador.nome)
            .all()
        )
    else:
        colaboradores = Colaborador.query.filter(Colaborador.id.in_(ids_colab)).order_by(Colaborador.nome).all()

    if departamentos_ids:
        departamentos_ids_set = set(int(x) for x in departamentos_ids if x is not None)
        colaboradores = [c for c in colaboradores if c.departamento_id and c.departamento_id in departamentos_ids_set]

    all_colab_ids = [c.id for c in colaboradores] or ids_colab
    assid_map = {}
    if all_colab_ids:
        for rec in PlrAssiduidade.query.filter(PlrAssiduidade.colaborador_id.in_(all_colab_ids)).all():
            assid_map[(rec.colaborador_id, rec.mes, rec.ano)] = assiduidade_pct_por_faltas(rec.faltas)

    av_por_colab_mes = {}
    for av in avaliacoes_periodo:
        cid = av.colaborador_id
        if cid not in av_por_colab_mes:
            av_por_colab_mes[cid] = {}
        mes_key = (av.data.month, av.data.year) if av.data else None
        if mes_key:
            assid_pct = assid_map.get((cid, mes_key[0], mes_key[1]))
            if assid_pct is not None:
                media = nota_media_com_assiduidade(av.avaliacao, assid_pct)
            else:
                media = av.nota_media_avaliacao()
            if media is not None:
                pct = min(100.0, max(0.0, float(media) * 10.0))
                if mes_key not in av_por_colab_mes[cid]:
                    av_por_colab_mes[cid][mes_key] = []
                av_por_colab_mes[cid][mes_key].append(pct)

    resultado = []
    for colab in colaboradores:
        if colab.data_admissao and colab.data_admissao > data_fechamento:
            continue
        data_ref = (
            colab.data_demissao
            if colab.data_demissao and colab.data_demissao <= data_fechamento
            else data_fechamento
        )
        tempo_meses = tempo_de_casa_meses(colab.data_admissao, data_ref)
        print("tempo_meses: ", colab.nome, tempo_meses, data_ref)
        if tempo_meses < 3:
            continue
        ref_mes, ref_ano = meses_colunas[-1] if meses_colunas else (mes_f, ano_f)
        dados_sal = salario_base_plr(colab, ref_mes, ref_ano, data_fechamento, db.session)
        salario_base_plr_val = dados_sal.get('salario_base_plr')

        pcts_meses = []
        soma_pct = 0.0
        for (m, a) in meses_colunas:
            mes_key = (m, a)
            listas_pct = av_por_colab_mes.get(colab.id, {}).get(mes_key, [])
            pct = sum(listas_pct) / len(listas_pct) if listas_pct else None
            pcts_meses.append(pct)
            if pct is not None:
                soma_pct += pct

        num_meses = len(meses_colunas)
        p_val = (soma_pct / num_meses) if num_meses and soma_pct is not None else None
        vpo = (salario_base_plr_val / 12.0) * 6 * (p_val / 100.0) if (
            salario_base_plr_val is not None and p_val is not None
        ) else None
        valor_total = vpo

        resultado.append({
            'colaborador': colab,
            'tempo_casa_meses': tempo_meses,
            'data_fechamento': data_fechamento,
            'salario_base_plr': salario_base_plr_val,
            'meses_colunas': meses_colunas,
            'pcts_meses': pcts_meses,
            'soma': round(soma_pct, 2) if soma_pct else None,
            'p': round(p_val, 2) if p_val is not None else None,
            'vpo': round(vpo, 2) if vpo is not None else None,
            'valor_total': round(valor_total, 2) if valor_total is not None else None,
        })

    resultado.sort(key=lambda x: (x['colaborador'].nome if x['colaborador'] else ''))
    return resultado, meses_colunas, data_fechamento


def _row_planilha_para_json(r):
    """Converte um item de resultado da planilha em dict para JSON (DataTables)."""
    colab = r['colaborador']
    data_fech = r.get('data_fechamento')
    demissao_ou_fech = (
        colab.data_demissao.strftime('%d/%m/%Y')
        if colab.data_demissao
        else (data_fech.strftime('%d/%m/%Y') if data_fech else '-')
    )
    return {
        'cpf': colab.cpf or '-',
        'nome': (colab.nome or '').upper(),
        'admissao': colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else '-',
        'demissao_fechamento': demissao_ou_fech,
        'tempo_casa_meses': r['tempo_casa_meses'],
        'funcao': colab.cargo.nome if colab.cargo else '-',
        'pcts_meses': r['pcts_meses'],
        'soma': f"{r['soma']:.2f}%" if r['soma'] is not None else '-',
        'p': f"{r['p']:.2f}%" if r['p'] is not None else '-',
        'salario_base_plr': (
            f"R$ {r['salario_base_plr']:.2f}".replace('.', ',')
            if r['salario_base_plr'] is not None else '-'
        ),
        'vpo': f"R$ {r['vpo']:.2f}".replace('.', ',') if r['vpo'] is not None else '-',
        'valor_total': (
            f"R$ {r['valor_total']:.2f}".replace('.', ',')
            if r['valor_total'] is not None else '-'
        ),
    }


# ---------- Rotas ----------

@plr_bp.route('/relatorio-planilha/dados')
def relatorio_planilha_dados():
    """Retorna JSON com os dados da planilha para DataTables (AJAX)."""
    ano_inicio = request.args.get('ano_inicio') or request.args.get('ano')
    ano_fim = request.args.get('ano_fim') or ano_inicio
    mes_inicio = request.args.get('mes_inicio', '1')
    mes_fim = request.args.get('mes_fim', '12')
    modelo_plr_id = request.args.get('modelo_plr_id', '')
    equipe_filtro = request.args.get('equipe', '').strip() or None
    departamentos_ids = [int(x) for x in request.args.getlist('departamento_id') if x and str(x).isdigit()]

    if not ano_inicio:
        return jsonify({
            'error': 'Informe o período.',
            'data': [], 'meses_colunas': [], 'data_fechamento': None
        }), 400
    try:
        ano_inicio = int(ano_inicio)
        ano_fim = int(ano_fim)
    except (ValueError, TypeError):
        return jsonify({
            'error': 'Ano inválido.',
            'data': [], 'meses_colunas': [], 'data_fechamento': None
        }), 400
    try:
        resultado, meses_colunas, data_fechamento = relatorio_planilha_calcular(
            ano_inicio, mes_inicio, ano_fim, mes_fim,
            modelo_plr_id, equipe_filtro, departamentos_ids or None
        )
    except Exception as e:
        return jsonify({
            'error': str(e),
            'data': [], 'meses_colunas': [], 'data_fechamento': None
        }), 500

    data = [_row_planilha_para_json(r) for r in resultado]
    return jsonify({
        'data': data,
        'meses_colunas': [{'mes': m, 'ano': a} for (m, a) in meses_colunas],
        'data_fechamento': data_fechamento.strftime('%d/%m/%Y') if data_fechamento else None,
    })


@plr_bp.route('/relatorio-planilha/excel')
def relatorio_planilha_excel():
    """Gera arquivo Excel com aba resumo e uma aba por mês, usando os mesmos filtros do relatório."""
    ano_inicio = request.args.get('ano_inicio') or request.args.get('ano')
    ano_fim = request.args.get('ano_fim') or ano_inicio
    mes_inicio = request.args.get('mes_inicio', '1')
    mes_fim = request.args.get('mes_fim', '12')
    modelo_plr_id = request.args.get('modelo_plr_id', '')
    equipe_filtro = request.args.get('equipe', '').strip() or None
    departamentos_ids = [int(x) for x in request.args.getlist('departamento_id') if x and str(x).isdigit()]

    if not ano_inicio:
        flash('Informe o período para exportar o Excel.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        ano_inicio_int = int(ano_inicio)
        ano_fim_int = int(ano_fim)
    except (ValueError, TypeError):
        flash('Ano inválido para exportação do Excel.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        resultado, meses_colunas, data_fechamento = relatorio_planilha_calcular(
            ano_inicio_int, mes_inicio, ano_fim_int, mes_fim,
            modelo_plr_id, equipe_filtro, departamentos_ids or None
        )
    except Exception as e:
        flash(f'Erro ao gerar Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    data_inicio = date(int(ano_inicio), int(mes_inicio), 1)
    cpf_para_chapa = _excel_cpf_para_chapa(resultado, data_inicio, data_fechamento)
    criterios_por_colab_mes = _excel_criterios_por_colab_mes(
        resultado, meses_colunas, data_inicio, data_fechamento,
        modelo_plr_id, equipe_filtro
    )

    wb = Workbook()
    estilos = _excel_estilos()
    _excel_aba_resumo(wb, resultado, meses_colunas, cpf_para_chapa, estilos)
    _excel_abas_meses(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Planilha_{ano_inicio}-{mes_inicio}_a_{ano_fim}-{mes_fim}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@plr_bp.route('/relatorio-planilha/excel-formulas')
def relatorio_planilha_excel_formulas():
    """
    Gera o mesmo Excel do relatório planilha, porém com SOMA, P, VPO e Valor total
    calculados por fórmulas do Excel (e totalizador em fórmula).
    """
    ano_inicio = request.args.get('ano_inicio') or request.args.get('ano')
    ano_fim = request.args.get('ano_fim') or ano_inicio
    mes_inicio = request.args.get('mes_inicio', '1')
    mes_fim = request.args.get('mes_fim', '12')
    modelo_plr_id = request.args.get('modelo_plr_id', '')
    equipe_filtro = request.args.get('equipe', '').strip() or None
    departamentos_ids = [int(x) for x in request.args.getlist('departamento_id') if x and str(x).isdigit()]

    if not ano_inicio:
        flash('Informe o período para exportar o Excel.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        ano_inicio_int = int(ano_inicio)
        ano_fim_int = int(ano_fim)
    except (ValueError, TypeError):
        flash('Ano inválido para exportação do Excel.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        resultado, meses_colunas, data_fechamento = relatorio_planilha_calcular(
            ano_inicio_int, mes_inicio, ano_fim_int, mes_fim,
            modelo_plr_id, equipe_filtro, departamentos_ids or None
        )
    except Exception as e:
        flash(f'Erro ao gerar Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    data_inicio = date(int(ano_inicio), int(mes_inicio), 1)
    cpf_para_chapa = _excel_cpf_para_chapa(resultado, data_inicio, data_fechamento)
    criterios_por_colab_mes = _excel_criterios_por_colab_mes(
        resultado, meses_colunas, data_inicio, data_fechamento,
        modelo_plr_id, equipe_filtro
    )

    wb = Workbook()
    wb.remove(wb.worksheets[0])
    estilos = _excel_estilos()
    _excel_abas_meses_formulas(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos)
    _excel_aba_resumo_formulas(wb, resultado, meses_colunas, cpf_para_chapa, estilos)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Planilha_Formulas_{ano_inicio}-{mes_inicio}_a_{ano_fim}-{mes_fim}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


def _assiduidade_parse_filtros():
    """Extrai filtros da request (GET ou form) para período/modelo/equipe/departamentos."""
    ano_inicio = request.args.get('ano_inicio') or request.form.get('ano_inicio') or request.args.get('ano')
    ano_fim = request.args.get('ano_fim') or request.form.get('ano_fim') or ano_inicio
    mes_inicio = request.args.get('mes_inicio') or request.form.get('mes_inicio', '1')
    mes_fim = request.args.get('mes_fim') or request.form.get('mes_fim', '12')
    modelo_plr_id = request.args.get('modelo_plr_id') or request.form.get('modelo_plr_id', '')
    equipe_filtro = (request.args.get('equipe') or request.form.get('equipe') or '').strip() or None
    dep_ids = request.args.getlist('departamento_id') or request.form.getlist('departamento_id')
    departamentos_ids = [int(x) for x in dep_ids if x and str(x).isdigit()]
    return ano_inicio, ano_fim, mes_inicio, mes_fim, modelo_plr_id, equipe_filtro, departamentos_ids


@plr_bp.route('/relatorio-planilha/assiduidade/template')
def relatorio_planilha_assiduidade_template():
    """Gera Excel template para importação de assiduidade: linhas = colaboradores (CPF, Nome), colunas = meses (faltas)."""
    ano_inicio, ano_fim, mes_inicio, mes_fim, modelo_plr_id, equipe_filtro, departamentos_ids = _assiduidade_parse_filtros()
    if not ano_inicio:
        flash('Informe o período (ano_inicio) para gerar o template.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        ano_inicio = int(ano_inicio)
        ano_fim = int(ano_fim)
    except (ValueError, TypeError):
        flash('Ano inválido.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    try:
        resultado, meses_colunas, _ = relatorio_planilha_calcular(
            ano_inicio, mes_inicio, ano_fim, mes_fim,
            modelo_plr_id, equipe_filtro, departamentos_ids or None
        )
    except Exception as e:
        flash(f'Erro ao gerar template: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Assiduidade'
    header = ['CPF', 'Nome']
    for (mes, ano) in meses_colunas:
        idx = int(mes) - 1
        mes_nome = MESES_ABREV[idx] if 0 <= idx < len(MESES_ABREV) else f'{mes:02d}'
        header.append(f'{mes_nome} {ano}')
    ws.append(header)
    for r in resultado:
        colab = r.get('colaborador')
        if not colab:
            continue
        cpf = colab.cpf or ''
        nome = (colab.nome or '').upper()
        row = [cpf, nome]
        for _ in meses_colunas:
            row.append(0)
        ws.append(row)
    estilos = _excel_estilos()
    for cell in ws[1]:
        cell.font = estilos['header_font']
        cell.fill = estilos['header_fill']
        cell.alignment = estilos['center_align']
    for col_idx, _ in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14 if col_idx <= 2 else 7.3
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'PLR_Assiduidade_Template_{ano_inicio}-{mes_inicio}_a_{ano_fim}-{mes_fim}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


def _parse_meses_do_header(header_row):
    """
    Lê os cabeçalhos da planilha (a partir da coluna 3) e retorna lista [(mes, ano), ...]
    extraídos de rótulos como 'JAN 2025', 'FEV 2025', 'MAR/2025', '01/2025', etc.
    """
    abrev_to_num = {abrev: i + 1 for i, abrev in enumerate(MESES_ABREV)}
    meses = []
    for cell_val in (header_row[2:] if len(header_row) > 2 else []):
        if cell_val is None:
            continue
        txt = str(cell_val).strip().upper()
        if not txt:
            continue
        parts = txt.replace('/', ' ').replace('-', ' ').split()
        if len(parts) == 2:
            label, ano_str = parts[0], parts[1]
            try:
                ano = int(ano_str)
            except ValueError:
                continue
            if label in abrev_to_num:
                meses.append((abrev_to_num[label], ano))
            else:
                try:
                    m = int(label)
                    if 1 <= m <= 12:
                        meses.append((m, ano))
                except ValueError:
                    continue
    return meses


def _normalizar_cpf(cpf):
    """Retorna CPF apenas com dígitos para comparação."""
    if cpf is None:
        return ''
    return ''.join(c for c in str(cpf) if c.isdigit())


def _mapa_cpf_colaborador():
    """Retorna dict cpf_normalizado -> Colaborador para todos os colaboradores."""
    colabs = Colaborador.query.all()
    return {_normalizar_cpf(c.cpf): c for c in colabs if _normalizar_cpf(c.cpf)}


@plr_bp.route('/relatorio-planilha/assiduidade/preview-meses', methods=['POST'])
def relatorio_planilha_assiduidade_preview_meses():
    """Retorna JSON com os meses/anos detectados no cabeçalho do arquivo Excel enviado."""
    arquivo = request.files.get('arquivo') or request.files.get('file')
    if not arquivo or not arquivo.filename:
        return jsonify({'ok': False, 'error': 'Nenhum arquivo enviado.'}), 400
    try:
        from openpyxl import load_workbook
        wb = load_workbook(arquivo, read_only=True, data_only=True)
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'Erro ao ler o Excel: {e}'}), 400
    if not header:
        return jsonify({'ok': False, 'error': 'Planilha sem cabeçalho.'}), 400

    meses = _parse_meses_do_header(list(header))
    if not meses:
        return jsonify({'ok': False, 'error': 'Nenhum mês/ano identificado nos cabeçalhos.'}), 400

    labels = [f'{MESES_ABREV[m - 1]}/{a}' for m, a in meses]
    return jsonify({'ok': True, 'meses': labels})


@plr_bp.route('/relatorio-planilha/assiduidade/import', methods=['POST'])
def relatorio_planilha_assiduidade_import():
    """Importa planilha Excel de assiduidade (faltas por colaborador por mês) e grava em PlrAssiduidade.
    Os meses/anos são detectados automaticamente a partir dos cabeçalhos do arquivo."""
    if 'arquivo' not in request.files and 'file' not in request.files:
        flash('Nenhum arquivo enviado.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))
    arquivo = request.files.get('arquivo') or request.files.get('file')
    if not arquivo or not arquivo.filename or not arquivo.filename.lower().endswith(('.xlsx', '.xls')):
        flash('Envie um arquivo Excel (.xlsx).', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    try:
        from openpyxl import load_workbook
        wb = load_workbook(arquivo, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        flash(f'Erro ao ler o Excel: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    if len(rows) < 2:
        flash('Planilha deve ter cabeçalho e ao menos uma linha de dados.', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    header = rows[0]
    meses_colunas = _parse_meses_do_header(header)

    if not meses_colunas:
        flash(
            'Não foi possível identificar os meses/ano nos cabeçalhos da planilha. '
            'Use o formato do template (ex: JAN 2025, FEV 2025).',
            'danger',
        )
        return redirect(url_for('plr.relatorio_planilha'))

    meses_label = ', '.join(
        f'{MESES_ABREV[m - 1]}/{a}' for m, a in meses_colunas
    )

    colunas_meses = len(meses_colunas)
    inseridos = 0
    atualizados = 0
    erros = []
    cpf_para_colab = _mapa_cpf_colaborador()

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or len(row) < 2:
            continue
        cpf_cel = row[0]
        cpf_limpo = _normalizar_cpf(cpf_cel)
        if not cpf_limpo or len(cpf_limpo) < 11:
            continue
        colab = cpf_para_colab.get(cpf_limpo)
        if not colab:
            erros.append(f'Linha {row_idx}: CPF não encontrado ({cpf_cel})')
            continue
        for col_idx in range(colunas_meses):
            if col_idx + 2 >= len(row):
                break
            mes, ano = meses_colunas[col_idx]
            val = row[col_idx + 2]
            try:
                faltas = int(float(val)) if val is not None else 0
            except (TypeError, ValueError):
                faltas = 0
            faltas = max(0, faltas)
            rec = PlrAssiduidade.query.filter_by(
                colaborador_id=colab.id, mes=mes, ano=ano
            ).first()
            if rec:
                rec.faltas = faltas
                atualizados += 1
            else:
                db.session.add(PlrAssiduidade(colaborador_id=colab.id, mes=mes, ano=ano, faltas=faltas))
                inseridos += 1
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao salvar: {e}', 'danger')
        return redirect(url_for('plr.relatorio_planilha'))

    msg_periodo = f' Período detectado: {meses_label}.'
    if erros:
        flash(f'Importado: {inseridos} novos, {atualizados} atualizados.{msg_periodo} Avisos: ' + '; '.join(erros[:5]), 'warning')
    else:
        flash(f'Assiduidade importada: {inseridos} novos, {atualizados} atualizados.{msg_periodo}', 'success')
    next_url = (request.form.get('next') or '').strip()
    if next_url and next_url.startswith('/'):
        return redirect(next_url)
    return redirect(url_for('plr.relatorio_planilha'))


@plr_bp.route('/relatorio-planilha/', methods=['GET', 'POST'])
def relatorio_planilha():
    """Página do relatório planilha: filtros e tabela (dados via AJAX ou POST)."""
    from datetime import datetime
    from models.plr import ModeloPLR

    modelos_list = ModeloPLR.query.filter_by(ativo=True).order_by(ModeloPLR.nome).all()
    equipes = _equipes_distintas_plr()
    departamentos = Departamento.query.filter_by(status='Ativo').order_by(Departamento.nome).all()
    ano_sugerido = datetime.now().year

    if request.method == 'POST':
        modelo_plr_id = request.form.get('modelo_plr_id', '')
        equipe_filtro = request.form.get('equipe', '').strip() or None
        departamentos_ids = [int(x) for x in request.form.getlist('departamento_id') if x and str(x).isdigit()]
        periodo_inicio = request.form.get('periodo_inicio')
        periodo_fim = request.form.get('periodo_fim')

        if periodo_inicio and periodo_fim:
            parts_i = periodo_inicio.split('-')
            parts_f = periodo_fim.split('-')
            ano_inicio = int(parts_i[0])
            mes_inicio = parts_i[1].lstrip('0') or '1'
            ano_fim = int(parts_f[0])
            mes_fim = parts_f[1].lstrip('0') or '12'
            if ano_inicio > ano_fim or (ano_inicio == ano_fim and int(mes_inicio) > int(mes_fim)):
                flash('A data de início deve ser anterior ou igual à data de fim.', 'danger')
                return _render_planilha(modelos_list, equipes, departamentos, ano_sugerido)
        else:
            ano_inicio = request.form.get('ano_inicio') or request.form.get('ano')
            ano_fim = request.form.get('ano_fim') or ano_inicio
            mes_inicio = request.form.get('mes_inicio', '1')
            mes_fim = request.form.get('mes_fim', '12')
            if not ano_inicio:
                flash('Informe o período.', 'danger')
                return _render_planilha(modelos_list, equipes, departamentos, ano_sugerido)
            ano_inicio = int(ano_inicio)
            ano_fim = int(ano_fim)

        try:
            resultado, meses_colunas, data_fechamento = relatorio_planilha_calcular(
                ano_inicio, mes_inicio, ano_fim, mes_fim,
                modelo_plr_id, equipe_filtro, departamentos_ids or None
            )
        except Exception as e:
            flash(f'Erro ao gerar relatório: {e}', 'danger')
            return _render_planilha(modelos_list, equipes, departamentos, ano_sugerido)

        planilha_json = {
            'data': [_row_planilha_para_json(r) for r in resultado],
            'meses_colunas': [{'mes': m, 'ano': a} for (m, a) in meses_colunas],
            'data_fechamento': data_fechamento.strftime('%d/%m/%Y') if data_fechamento else None,
        }
        return _render_planilha(
            modelos_list, equipes, departamentos, ano_sugerido,
            planilha_json=planilha_json
        )

    return _render_planilha(modelos_list, equipes, departamentos, ano_sugerido, planilha_json=None)


def _render_planilha(modelos_list, equipes, departamentos, ano_sugerido, planilha_json=None):
    return render_template(
        'plr/relatorio_planilha.html',
        modelos=modelos_list,
        equipes=equipes,
        departamentos=departamentos,
        ano_sugerido=ano_sugerido,
        planilha_json=planilha_json,
    )


# ---------- Helpers Excel ----------

def _excel_cpf_para_chapa(resultado, data_inicio, data_fechamento):
    """
    Mapa CPF -> chapa a partir do EfetivoPLR no período.
    Usa o último efetivo (por data) que tiver chapa disponível para cada CPF.
    """
    cpf_set = set()
    for r in resultado:
        colab = r.get('colaborador')
        if colab and colab.cpf:
            cpf_set.add(colab.cpf)
    cpf_para_chapa = {}
    if cpf_set:
        efetivos = (
            EfetivoPLR.query
            .filter(EfetivoPLR.cpf.in_(cpf_set))
            .filter(EfetivoPLR.data >= data_inicio, EfetivoPLR.data <= data_fechamento)
            .order_by(EfetivoPLR.cpf.asc(), EfetivoPLR.data.desc())
            .all()
        )
        for ef in efetivos:
            if not ef.cpf:
                continue
            chapa_val = (ef.chapa or '').strip()
            if ef.cpf not in cpf_para_chapa:
                cpf_para_chapa[ef.cpf] = chapa_val
            elif chapa_val and not cpf_para_chapa.get(ef.cpf):
                # Já tinha vazio; pega chapa de registro mais antigo que tenha chapa
                cpf_para_chapa[ef.cpf] = chapa_val
    return cpf_para_chapa


def _excel_criterios_por_colab_mes(resultado, meses_colunas, data_inicio, data_fechamento,
                                   modelo_plr_id, equipe_filtro):
    """Mapa (colaborador_id, (mes, ano)) -> dict de critérios (média por tipo)."""
    criterios_por_colab_mes = {}
    if not resultado or not meses_colunas:
        return criterios_por_colab_mes
    colab_ids = [r['colaborador'].id for r in resultado if r.get('colaborador')]
    if not colab_ids:
        return criterios_por_colab_mes

    q_crit = PLRColaborador.query.filter(
        PLRColaborador.colaborador_id.in_(colab_ids),
        PLRColaborador.data >= data_inicio,
        PLRColaborador.data <= data_fechamento,
    )
    if modelo_plr_id:
        q_crit = q_crit.filter(PLRColaborador.PlrModelo_id == int(modelo_plr_id))
    avaliacoes_crit = q_crit.all()
    if equipe_filtro:
        avaliacoes_crit = [
            av for av in avaliacoes_crit
            if av.equipe_alocada and isinstance(av.equipe_alocada, list)
            and equipe_filtro in [str(e).strip() for e in av.equipe_alocada if e]
        ]

    acumulado = {}
    for av in avaliacoes_crit:
        if not av.data:
            continue
        cid = av.colaborador_id
        mes_key = (av.data.month, av.data.year)
        key = (cid, mes_key)
        if key not in acumulado:
            acumulado[key] = {
                'Assiduidade': [],
                'Zero Acidente': [],
                'Segurança, Limpeza, Organização': [],
                'Prazo': [],
            }
        for tipo in ['Assiduidade', 'Zero Acidente', 'Segurança, Limpeza, Organização', 'Prazo']:
            v = _valor_por_tipo(av.avaliacao, tipo)
            if v is not None:
                acumulado[key][tipo].append(v)

    for (cid, mes_key), valores in acumulado.items():
        if cid not in criterios_por_colab_mes:
            criterios_por_colab_mes[cid] = {}
        criterios_por_colab_mes[cid][mes_key] = {
            tipo: (sum(lst) / len(lst) if lst else None)
            for tipo, lst in valores.items()
        }

    # Sobrescrever Assiduidade com dados de PlrAssiduidade quando existir (faltas -> %)
    meses_set = set(meses_colunas)
    assid_records = PlrAssiduidade.query.filter(
        PlrAssiduidade.colaborador_id.in_(colab_ids),
    ).all()
    for rec in assid_records:
        mes_key = (rec.mes, rec.ano)
        if mes_key not in meses_set:
            continue
        pct = assiduidade_pct_por_faltas(rec.faltas)
        if rec.colaborador_id not in criterios_por_colab_mes:
            criterios_por_colab_mes[rec.colaborador_id] = {}
        if mes_key not in criterios_por_colab_mes[rec.colaborador_id]:
            criterios_por_colab_mes[rec.colaborador_id][mes_key] = {}
        criterios_por_colab_mes[rec.colaborador_id][mes_key]['Assiduidade'] = pct

    return criterios_por_colab_mes


def _excel_estilos():
    """Estilos reutilizáveis para cabeçalho e células."""
    return {
        'header_font': Font(bold=True, color="FFFFFF"),
        'header_fill': PatternFill("solid", fgColor="4F81BD"),
        'center_align': Alignment(horizontal="center", vertical="center", wrap_text=True),
        'right_align': Alignment(horizontal="right", vertical="center"),
        'thin_border': Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        ),
    }


def _excel_aba_resumo(wb, resultado, meses_colunas, cpf_para_chapa, estilos):
    """Preenche e formata a aba Resumo."""
    ws = wb.active
    ws.title = 'Resumo'
    header = ['CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão/Fech.', 'Tempo casa (meses)', 'Função']
    for (mes, ano) in meses_colunas:
        header.append(f'{_titulo_aba_mes(mes, ano)}')
    header.extend(['SOMA (%)', 'P (%)', 'Salário base PLR', 'VPO', 'Valor total'])
    ws.append(header)

    total_geral = 0
    for r in resultado:
        pcts_meses = r.get('pcts_meses') or []
        meses_com_avaliacao = [pct for pct in pcts_meses if pct is not None]
        if len(meses_com_avaliacao) < MIN_MESES_AVALIACAO_RESUMO:
            continue
        total_geral += (r.get('valor_total') or 0)
        colab = r['colaborador']
        data_fech = r.get('data_fechamento')
        cpf = colab.cpf or '-'
        chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
        admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else '-'
        if colab.data_demissao and data_fech and colab.data_demissao <= data_fech:
            demissao_fech = colab.data_demissao.strftime('%d/%m/%Y')
        else:
            demissao_fech = data_fech.strftime('%d/%m/%Y') if data_fech else '-'
        funcao = colab.cargo.nome if colab.cargo else '-'
        valores_mensais = [
            (pcts_meses[i] / 100.0) if i < len(pcts_meses) and pcts_meses[i] is not None else None
            for i in range(len(meses_colunas))
        ]
        linha = (
            [cpf, chapa, (colab.nome or '').upper(), admissao, demissao_fech, r['tempo_casa_meses'], funcao]
            + valores_mensais
            + [
                (r['soma'] / 100.0) if r['soma'] is not None else None,
                (r['p'] / 100.0) if r['p'] is not None else None,
                r['salario_base_plr'], r['vpo'], r['valor_total'],
            ]
        )
        ws.append(linha)

    # Linha de total: mescla células e exibe totalizador do valor total
    num_cols = len(header)
    total_row = ws.max_row + 1
    row_data = ['TOTAL'] + [None] * (num_cols - 2) + [total_geral]
    ws.append(row_data)
    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=num_cols - 1,
    )
    cell_total_label = ws.cell(row=total_row, column=1)
    cell_total_label.value = 'TOTAL'
    cell_total_label.font = Font(bold=True)
    cell_total_label.alignment = estilos['right_align']
    cell_total_label.border = estilos['thin_border']
    cell_total_val = ws.cell(row=total_row, column=num_cols)
    cell_total_val.number_format = r'R$ #,##0.00'
    cell_total_val.font = Font(bold=True)
    cell_total_val.alignment = estilos['right_align']
    cell_total_val.border = estilos['thin_border']
    for col in range(2, num_cols):
        c = ws.cell(row=total_row, column=col)
        c.border = estilos['thin_border']

    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = estilos['header_font']
            cell.fill = estilos['header_fill']
            cell.alignment = estilos['center_align']
            cell.border = estilos['thin_border']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = estilos['thin_border']

    num_meses = len(meses_colunas)
    col_mes_inicio = 8
    col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 12, 'G': 29}
    for idx_col in range(col_mes_inicio, col_mes_inicio + num_meses):
        col_letter = ws.cell(row=1, column=idx_col).column_letter
        col_widths[col_letter] = 7.3
    for idx_col in range(col_mes_inicio + num_meses, ws.max_column + 1):
        col_letter = ws.cell(row=1, column=idx_col).column_letter
        col_widths[col_letter] = 14
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    right = estilos['right_align']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = right
    if num_meses:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_mes_inicio, max_col=col_mes_inicio + num_meses - 1):
            for cell in row:
                cell.number_format = "0.00%"
                cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
        for cell in col:
            cell.number_format = "0"
            cell.alignment = right
    soma_col = col_mes_inicio + num_meses + 1
    p_col = soma_col + 1
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=soma_col, max_col=p_col):
        for cell in col:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col_idx in (p_col + 1, p_col + 2, p_col + 3):
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in col:
                cell.number_format = r'R$ #,##0.00'
                cell.alignment = right


def _excel_aba_resumo_formulas(wb, resultado, meses_colunas, cpf_para_chapa, estilos):
    """
    Preenche a aba Resumo inserida no início do workbook. O % de cada mês vem da aba
    do mês correspondente (fórmula INDEX/MATCH por CPF). SOMA, P, VPO e Valor total em fórmulas.
    """
    ws = wb.create_sheet('Resumo', 0)
    num_meses = len(meses_colunas)
    col_mes_inicio = 8
    col_mes_fim = 7 + num_meses
    col_soma = 8 + num_meses
    col_p = 9 + num_meses
    col_salario = 10 + num_meses
    col_vpo = 11 + num_meses
    col_valor_total = 12 + num_meses
    col_mes_pct_na_aba_mes = 11

    header = ['CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão/Fech.', 'Tempo casa (meses)', 'Função']
    for (mes, ano) in meses_colunas:
        header.append(_titulo_aba_mes(mes, ano))
    header.extend(['SOMA (%)', 'P (%)', 'Salário base PLR', 'VPO', 'Valor total'])
    ws.append(header)

    range_mes_ini = get_column_letter(col_mes_inicio)
    range_mes_fim = get_column_letter(col_mes_fim)
    letter_salario = get_column_letter(col_salario)
    letter_p = get_column_letter(col_p)
    letter_vpo = get_column_letter(col_vpo)
    letter_valor_total = get_column_letter(col_valor_total)
    col_mes_letter_aba = get_column_letter(col_mes_pct_na_aba_mes)

    data_row = 1
    for r in resultado:
        if False: 
            pcts_meses = r.get('pcts_meses') or []
            meses_com_avaliacao = [pct for pct in pcts_meses if pct is not None]
            if len(meses_com_avaliacao) < MIN_MESES_AVALIACAO_RESUMO:
                continue
        data_row += 1
        row_idx = data_row
        colab = r['colaborador']
        data_fech = r.get('data_fechamento')
        cpf = colab.cpf or '-'
        chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
        admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else '-'
        if colab.data_demissao and data_fech and colab.data_demissao <= data_fech:
            demissao_fech = colab.data_demissao.strftime('%d/%m/%Y')
        else:
            demissao_fech = data_fech.strftime('%d/%m/%Y') if data_fech else '-'
        funcao = colab.cargo.nome if colab.cargo else '-'
        for col, val in enumerate([cpf, chapa, (colab.nome or '').upper(), admissao, demissao_fech, r['tempo_casa_meses'], funcao], start=1):
            ws.cell(row=row_idx, column=col, value=val)
        for i, (mes, ano) in enumerate(meses_colunas):
            sheet_mes = f"'{_titulo_aba_mes(mes, ano)}'"
            col_resumo = col_mes_inicio + i
            ws.cell(
                row=row_idx,
                column=col_resumo,
                value=f"=IFERROR(INDEX({sheet_mes}!${col_mes_letter_aba}:${col_mes_letter_aba},MATCH($A{row_idx},{sheet_mes}!$A:$A,0)),\"\")"
            )
        ws.cell(row=row_idx, column=col_salario, value=r.get('salario_base_plr'))
        ws.cell(row=row_idx, column=col_soma, value=f'=SUM({range_mes_ini}{row_idx}:{range_mes_fim}{row_idx})')
        ws.cell(row=row_idx, column=col_p, value=f'=SUM({range_mes_ini}{row_idx}:{range_mes_fim}{row_idx})/{num_meses}')
        ws.cell(row=row_idx, column=col_vpo, value=f'=({letter_salario}{row_idx}/12)*6')
        ws.cell(row=row_idx, column=col_valor_total, value=f'={letter_vpo}{row_idx}*{letter_p}{row_idx}')

    last_data_row = ws.max_row
    total_row = last_data_row + 1
    num_cols = col_valor_total
    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=num_cols - 1,
    )
    cell_total_label = ws.cell(row=total_row, column=1)
    cell_total_label.value = 'TOTAL'
    cell_total_label.font = Font(bold=True)
    cell_total_label.alignment = estilos['right_align']
    cell_total_label.border = estilos['thin_border']
    cell_total_val = ws.cell(row=total_row, column=num_cols)
    cell_total_val.value = f'=SUM({letter_valor_total}2:{letter_valor_total}{last_data_row})'
    cell_total_val.number_format = r'R$ #,##0.00'
    cell_total_val.font = Font(bold=True)
    cell_total_val.alignment = estilos['right_align']
    cell_total_val.border = estilos['thin_border']
    for col in range(2, num_cols):
        ws.cell(row=total_row, column=col).border = estilos['thin_border']

    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = estilos['header_font']
        cell.fill = estilos['header_fill']
        cell.alignment = estilos['center_align']
        cell.border = estilos['thin_border']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = estilos['thin_border']

    num_meses_resumo = len(meses_colunas)
    col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 12, 'G': 29}
    for idx_col in range(col_mes_inicio, col_mes_inicio + num_meses_resumo):
        col_widths[get_column_letter(idx_col)] = 7.3
    for idx_col in range(col_mes_inicio + num_meses_resumo, ws.max_column + 1):
        col_widths[get_column_letter(idx_col)] = 14
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    right = estilos['right_align']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = right
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_mes_inicio, max_col=col_mes_fim):
        for cell in row:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
        for cell in col:
            cell.number_format = "0"
            cell.alignment = right
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_soma, max_col=col_p):
        for cell in col:
            cell.number_format = "0.00%"
            cell.alignment = right
    for col_idx in (col_salario, col_vpo, col_valor_total):
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in col:
                cell.number_format = r'R$ #,##0.00'
                cell.alignment = right


def _excel_abas_meses_formulas(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos):
    """
    Cria uma aba por mês com fórmulas: % Mês = soma ponderada dos critérios,
    Qtd Meses e Acum (%) calculados por fórmulas com referências cruzadas entre abas.
    O Resumo usa o % Mês destas abas (fórmula INDEX/MATCH por CPF).
    """
    header_font = estilos['header_font']
    header_fill = estilos['header_fill']
    center_align = estilos['center_align']
    right_align = estilos['right_align']
    thin_border = estilos['thin_border']
    col_pct_mes = 11
    col_qtd = 12
    col_acum = 13
    letter_pct_mes = get_column_letter(col_pct_mes)

    headers_criterios = [label for _, label in _criterios_headers_com_peso()]
    pesos = _pesos_por_tipo()

    titulos_abas = [_titulo_aba_mes(m, a) for m, a in meses_colunas]

    for idx, (mes, ano) in enumerate(meses_colunas):
        titulo = titulos_abas[idx]
        ws = wb.create_sheet(title=titulo)
        header_mes = [
            'CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão', 'Função',
        ] + headers_criterios + [
            '% Mês', 'Qtd Meses', 'Acum (%)',
        ]
        ws.append(header_mes)

        data_row = 1
        for r in resultado:
            pct_mes_list = r.get('pcts_meses') or []
            pct_mes = pct_mes_list[idx] if idx < len(pct_mes_list) else None
            if pct_mes is None:
                continue
            data_row += 1
            row_idx = data_row
            colab = r['colaborador']
            cpf = colab.cpf or '-'
            chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
            admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else ''
            demissao = colab.data_demissao.strftime('%d/%m/%Y') if colab.data_demissao else ''
            funcao = colab.cargo.nome if colab.cargo else '-'

            crit_colab = criterios_por_colab_mes.get(colab.id, {})
            crit_mes = crit_colab.get((mes, ano), {}) if crit_colab else {}
            assid = _valor_criterio_ponderado_excel(crit_mes.get('Assiduidade'), pesos.get('Assiduidade'))
            zero_ac = _valor_criterio_ponderado_excel(crit_mes.get('Zero Acidente'), pesos.get('Zero Acidente'))
            segur = _valor_criterio_ponderado_excel(crit_mes.get('Segurança, Limpeza, Organização'), pesos.get('Segurança, Limpeza, Organização'))
            prazo = _valor_criterio_ponderado_excel(crit_mes.get('Prazo'), pesos.get('Prazo'))

            for col, val in enumerate([cpf, chapa, (colab.nome or '').upper(), admissao, demissao, funcao, assid, zero_ac, segur, prazo], start=1):
                ws.cell(row=row_idx, column=col, value=val)

            ws.cell(row=row_idx, column=col_pct_mes,
                    value=f'=G{row_idx}+H{row_idx}+I{row_idx}+J{row_idx}')

            qtd_parts = []
            sum_parts = []
            for j in range(idx + 1):
                if j == idx:
                    qtd_parts.append(f'IF({letter_pct_mes}{row_idx}>0,1,0)')
                    sum_parts.append(f'{letter_pct_mes}{row_idx}')
                else:
                    tab_ref = f"'{titulos_abas[j]}'"
                    qtd_parts.append(
                        f'IF(ISNUMBER(MATCH($A{row_idx},{tab_ref}!$A:$A,0)),1,0)'
                    )
                    sum_parts.append(
                        f'IFERROR(INDEX({tab_ref}!${letter_pct_mes}:${letter_pct_mes},MATCH($A{row_idx},{tab_ref}!$A:$A,0)),0)'
                    )

            ws.cell(row=row_idx, column=col_qtd,
                    value=f'={"+".join(qtd_parts)}')
            num_meses_periodo = idx + 1
            ws.cell(row=row_idx, column=col_acum,
                    value=f'=({"+".join(sum_parts)})/{num_meses_periodo}')

        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border

        col_widths = {'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 29, 'G': 14, 'H': 14, 'I': 22, 'J': 12, 'K': 10, 'L': 10, 'M': 10}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        for col_idx in range(4, 6):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "dd/mm/yyyy"
                    cell.alignment = right_align
        for col_idx in range(7, 11):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.0%"
                    cell.alignment = right_align
        for col_idx in (col_pct_mes, col_acum):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.00%"
                    cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_qtd, max_col=col_qtd):
            for cell in col:
                cell.number_format = "0"
                cell.alignment = right_align


def _excel_abas_meses(wb, resultado, meses_colunas, cpf_para_chapa, criterios_por_colab_mes, estilos):
    """Cria uma aba por mês com colaboradores que têm avaliação no mês."""
    header_font = estilos['header_font']
    header_fill = estilos['header_fill']
    center_align = estilos['center_align']
    right_align = estilos['right_align']
    thin_border = estilos['thin_border']
    headers_criterios = [label for _, label in _criterios_headers_com_peso()]
    pesos = _pesos_por_tipo()

    for idx, (mes, ano) in enumerate(meses_colunas):
        titulo = _titulo_aba_mes(mes, ano)
        ws = wb.create_sheet(title=titulo)
        header_mes = [
            'CPF', 'Chapa', 'Nome', 'Admissão', 'Demissão', 'Função',
        ] + headers_criterios + [
            '% Mês', 'Qtd Meses', 'Acum (%)',
        ]
        ws.append(header_mes)

        for r in resultado:
            colab = r['colaborador']
            pct_mes_list = r.get('pcts_meses') or []
            pct_mes = pct_mes_list[idx] if idx < len(pct_mes_list) else None
            if pct_mes is None:
                continue
            valid_prev = [p for j, p in enumerate(pct_mes_list) if j <= idx and p is not None]
            qtd_meses = len(valid_prev)
            num_meses_periodo = idx + 1
            acum_pct = (sum(valid_prev) / num_meses_periodo) if num_meses_periodo else None

            cpf = colab.cpf or '-'
            chapa = cpf_para_chapa.get(cpf, '') if cpf != '-' else ''
            admissao = colab.data_admissao.strftime('%d/%m/%Y') if colab.data_admissao else ''
            demissao = colab.data_demissao.strftime('%d/%m/%Y') if colab.data_demissao else ''
            funcao = colab.cargo.nome if colab.cargo else '-'
            crit_colab = criterios_por_colab_mes.get(colab.id, {})
            crit_mes = crit_colab.get((mes, ano), {}) if crit_colab else {}
            assid = _valor_criterio_ponderado_excel(crit_mes.get('Assiduidade'), pesos.get('Assiduidade'))
            zero_ac = _valor_criterio_ponderado_excel(crit_mes.get('Zero Acidente'), pesos.get('Zero Acidente'))
            segur = _valor_criterio_ponderado_excel(crit_mes.get('Segurança, Limpeza, Organização'), pesos.get('Segurança, Limpeza, Organização'))
            prazo = _valor_criterio_ponderado_excel(crit_mes.get('Prazo'), pesos.get('Prazo'))

            ws.append([
                cpf, chapa, (colab.nome or '').upper(), admissao, demissao, funcao,
                assid, zero_ac, segur, prazo,
                (pct_mes / 100.0),
                qtd_meses,
                (acum_pct / 100.0) if acum_pct is not None else None,
            ])

        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border

        col_widths_mes = {
            'A': 14, 'B': 10, 'C': 36, 'D': 12, 'E': 12, 'F': 29,
            'G': 14, 'H': 14, 'I': 22, 'J': 12, 'K': 10, 'L': 10, 'M': 10,
        }
        for col, width in col_widths_mes.items():
            ws.column_dimensions[col].width = width

        for col_idx in range(4, 6):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "dd/mm/yyyy"
                    cell.alignment = right_align
        for col_idx in range(7, 11):
            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:
                    cell.number_format = "0.0%"
                    cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=11, max_col=11):
            for cell in col:
                cell.number_format = "0.00%"
                cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=12, max_col=12):
            for cell in col:
                cell.number_format = "0"
                cell.alignment = right_align
        for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=13, max_col=13):
            for cell in col:
                cell.number_format = "0.00%"
                cell.alignment = right_align
