from flask import Blueprint, render_template, redirect, url_for, flash, request, jsonify, send_file, make_response
from flask_login import login_required, current_user
from datetime import datetime, date
from models.usuario import Usuario
import pandas as pd
import os
import tempfile
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from io import BytesIO
try:
    from weasyprint import HTML, CSS
    from weasyprint.text.fonts import FontConfiguration
except ImportError:
    HTML = None
    CSS = None
    FontConfiguration = None
    # Adicionar um log ou aviso se WeasyPrint não estiver instalado
    print("AVISO: WeasyPrint não está instalado. A geração de PDF não funcionará.")
import zipfile
import json
from sqlalchemy.orm import joinedload

from models.database import db
from models.colaborador import Colaborador, DadosBancarios
from models.colaborador.utils.mudanca_funcao import (
    historico_mudanca_funcao_normalizado,
    normalizar_id_opcional,
)
from models.departamento import Departamento
from models.cargo import Cargo
from models.usuario import Usuario
from models.epi import EpiEntregas, Epi
from models.concreto import ConcretoUsinagens

colaborador_bp = Blueprint('colaborador', __name__)


def _mesclar_matricula_em_dados_adicionais(dados_adicionais_existentes, matricula_raw):
    """Atualiza a chave `matricula` no JSON de dados_adicionais, preservando demais chaves."""
    base = {}
    if dados_adicionais_existentes:
        try:
            parsed = (
                json.loads(dados_adicionais_existentes)
                if isinstance(dados_adicionais_existentes, str)
                else dados_adicionais_existentes
            )
            if isinstance(parsed, dict):
                base = dict(parsed)
        except (json.JSONDecodeError, TypeError):
            base = {}
    matricula = (matricula_raw or '').strip()
    if matricula:
        base['matricula'] = matricula
        base.pop('chapa', None)
    else:
        base.pop('matricula', None)
        base.pop('chapa', None)
    if not base:
        return None
    return json.dumps(base, ensure_ascii=False)


# Middleware para verificar se o usuário tem permissão
#@colaborador_bp.before_request
#@login_required
#def verificar_permissao():
#    if not current_user.is_admin:
#        flash('Acesso restrito. Você não tem permissão para acessar esta área.', 'danger')
#        return redirect(url_for('dashboard.index'))

@colaborador_bp.route('/')
@login_required
def index():
    """
    Lista todos os colaboradores
    """
    colaboradores = Colaborador.query.order_by(Colaborador.nome.asc()).all()
    departamentos = Departamento.query.filter_by(status='Ativo').all()
    cargos = Cargo.query.filter_by(status='Ativo').all()
    
    return render_template('cadastros/colaboradores/index.html', 
                          colaboradores=colaboradores,
                          departamentos=departamentos,
                          cargos=cargos)

@colaborador_bp.route('/novo', methods=['POST'])
@login_required
def novo():
    """
    Cria um novo colaborador
    """
    nome = request.form.get('nome')
    cpf = request.form.get('cpf')
    rg = request.form.get('rg')
    data_nascimento_str = request.form.get('data_nascimento')
    data_admissao_str = request.form.get('data_admissao')
    cargo_id = request.form.get('cargo')
    departamento_id = request.form.get('departamento')
    status = request.form.get('status', 'Ativo')
    telefone = request.form.get('telefone')
    email = request.form.get('email')
    endereco = request.form.get('endereco')
    observacoes = request.form.get('observacoes')
    matricula = request.form.get('matricula')
    
    # Validação básica
    if not nome or not data_admissao_str or not cargo_id or not departamento_id:
        flash('Por favor, preencha todos os campos obrigatórios.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    # Verifica se já existe um colaborador com o mesmo CPF (se informado)
    if cpf:
        colaborador_existente = Colaborador.query.filter_by(cpf=cpf).first()
        if colaborador_existente:
            flash('Já existe um colaborador com este CPF.', 'danger')
            return redirect(url_for('colaborador.index'))
    
    dados_adicionais_json = _mesclar_matricula_em_dados_adicionais(None, matricula)
    
    # Tratamento das datas
    data_nascimento = None
    if data_nascimento_str:
        try:
            data_nascimento = datetime.strptime(data_nascimento_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Formato de data de nascimento inválido.', 'danger')
            return redirect(url_for('colaborador.index'))
    
    data_admissao = None
    try:
        data_admissao = datetime.strptime(data_admissao_str, '%Y-%m-%d').date()
    except ValueError:
        flash('Formato de data de admissão inválido.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    # Cria o novo colaborador
    novo_colaborador = Colaborador(
        nome=nome,
        cpf=cpf,
        rg=rg,
        data_nascimento=data_nascimento,
        data_admissao=data_admissao,
        cargo_id=cargo_id,
        departamento_id=departamento_id,
        status=status,
        telefone=telefone,
        email=email,
        endereco=endereco,
        observacoes=observacoes,
        dados_adicionais=dados_adicionais_json
    )
    
    novo_colaborador.save()
    
    # Verifica se precisa criar um usuário para este colaborador
    criar_usuario = request.form.get('criar_usuario')
    if criar_usuario == 'sim':
        # Gera um email se não foi fornecido
        email_usuario = email or f"{nome.lower().replace(' ', '.')}@fortanks.com.br"
        
        # Verifica se já existe usuário com este email
        usuario_existente = Usuario.query.filter_by(email=email_usuario).first()
        if not usuario_existente:
            # Define uma senha padrão (deve ser alterada no primeiro login)
            senha_padrao = 'mudar123'
            
            # Busca o nome do cargo e departamento
            cargo = Cargo.query.get(cargo_id)
            departamento = Departamento.query.get(departamento_id)
            
            novo_usuario = Usuario(
                nome=nome,
                email=email_usuario,
                senha=Usuario.hash_password(senha_padrao),
                cargo=cargo.nome if cargo else None,
                departamento=departamento.nome if departamento else None
            )
            db.session.add(novo_usuario)
            db.session.commit()
            
            flash(f'Usuário criado para o colaborador. Email: {email_usuario} - Senha: {senha_padrao}', 'info')
    
    flash('Colaborador criado com sucesso.', 'success')
    return redirect(url_for('colaborador.index'))

@colaborador_bp.route('/editar/<int:id>', methods=['POST'])
@login_required
def editar(id):
    """
    Edita um colaborador existente
    """
    colaborador = Colaborador.query.get_or_404(id)
    
    nome = request.form.get('nome')
    cpf = request.form.get('cpf')
    rg = request.form.get('rg')
    data_nascimento_str = request.form.get('data_nascimento')
    data_admissao_str = request.form.get('data_admissao')
    data_demissao_str = request.form.get('data_demissao')
    cargo_id = request.form.get('cargo')
    departamento_id = request.form.get('departamento')
    status = request.form.get('status')
    telefone = request.form.get('telefone')
    email = request.form.get('email')
    endereco = request.form.get('endereco')
    observacoes = request.form.get('observacoes')
    matricula = request.form.get('matricula')
    
    # Validação básica
    if not nome or not data_admissao_str or not cargo_id or not departamento_id:
        flash('Por favor, preencha todos os campos obrigatórios.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    # Verifica se já existe outro colaborador com o mesmo CPF (se informado)
    if cpf:
        colaborador_existente = Colaborador.query.filter_by(cpf=cpf).first()
        if colaborador_existente and colaborador_existente.id != id:
            flash('Já existe um colaborador com este CPF.', 'danger')
            return redirect(url_for('colaborador.index'))
    
    # Tratamento das datas
    data_nascimento = None
    if data_nascimento_str:
        try:
            data_nascimento = datetime.strptime(data_nascimento_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Formato de data de nascimento inválido.', 'danger')
            return redirect(url_for('colaborador.index'))
    
    data_admissao = None
    try:
        data_admissao = datetime.strptime(data_admissao_str, '%Y-%m-%d').date()
    except ValueError:
        flash('Formato de data de admissão inválido.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    data_demissao = None
    if data_demissao_str:
        try:
            data_demissao = datetime.strptime(data_demissao_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Formato de data de demissão inválido.', 'danger')
            return redirect(url_for('colaborador.index'))
    
    # Atualiza o colaborador
    colaborador.nome = nome
    colaborador.cpf = cpf
    colaborador.rg = rg
    colaborador.data_nascimento = data_nascimento
    colaborador.data_admissao = data_admissao
    colaborador.data_demissao = data_demissao
    colaborador.cargo_id = cargo_id
    colaborador.departamento_id = departamento_id
    colaborador.status = status
    colaborador.telefone = telefone
    colaborador.email = email
    colaborador.endereco = endereco
    colaborador.observacoes = observacoes
    colaborador.dados_adicionais = _mesclar_matricula_em_dados_adicionais(
        colaborador.dados_adicionais, matricula
    )
    colaborador.atualizado_em = datetime.now()
    
    colaborador.save()
    
    # Se mudou o status para "Inativo" ou "Demitido", verifica se precisa atualizar a data de demissão
    if status in ['Inativo', 'Demitido'] and not colaborador.data_demissao:
        colaborador.data_demissao = date.today()
        colaborador.save()
    
    flash('Colaborador atualizado com sucesso.', 'success')
    return redirect(url_for('colaborador.index'))

@colaborador_bp.route('/excluir/<int:id>', methods=['POST'])
@login_required
def excluir(id):
    """
    Exclui um colaborador
    """
    colaborador = Colaborador.query.get_or_404(id)
    
    nome_colaborador = colaborador.nome
    
    # Verifica se o colaborador tem vínculos que impedem a exclusão
    
    # Verificar se o colaborador possui entregas de EPIs
    entregas_epi = EpiEntregas.query.filter_by(colaborador_id=id).count()
    if entregas_epi > 0:
        flash(f'Não é possível excluir o colaborador {nome_colaborador} pois existem {entregas_epi} entregas de EPIs vinculadas a ele.', 'danger')
        return redirect(url_for('colaborador.index'))
        
    # Verificar se o colaborador é responsável por usinagens de concreto
    usinagens = ConcretoUsinagens.query.filter_by(responsavel_id=id).count()
    if usinagens > 0:
        flash(f'Não é possível excluir o colaborador {nome_colaborador} pois é responsável por {usinagens} usinagens de concreto.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    try:
        colaborador.delete()
        flash(f'Colaborador {nome_colaborador} excluído com sucesso.', 'success')
    except Exception as e:
        flash(f'Erro ao excluir o colaborador {nome_colaborador}: {str(e)}', 'danger')
    
    return redirect(url_for('colaborador.index'))

@colaborador_bp.route('/visualizar/<int:id>')
@login_required
def visualizar(id):
    """
    Visualiza os detalhes de um colaborador
    """
    colaborador = Colaborador.query.get_or_404(id)
    departamentos = Departamento.query.filter_by(status='Ativo').all()
    cargos = Cargo.query.filter_by(status='Ativo').all()
    
    return render_template('cadastros/colaboradores/visualizar.html', 
                          colaborador=colaborador,
                          departamentos=departamentos,
                          cargos=cargos)

@colaborador_bp.route('/importar_excel', methods=['POST'])
@login_required
def importar_excel():
    """
    Importa colaboradores de um arquivo Excel
    """
    if 'arquivo_excel' not in request.files:
        flash('Nenhum arquivo enviado.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    arquivo = request.files['arquivo_excel']
    
    if arquivo.filename == '':
        flash('Nenhum arquivo selecionado.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    if not arquivo.filename.endswith(('.xlsx', '.xls')):
        flash('Formato de arquivo inválido. Use arquivos Excel (.xlsx ou .xls).', 'danger')
        return redirect(url_for('colaborador.index'))
    
    try:
        # Salva o arquivo temporariamente
        filename = secure_filename(arquivo.filename)
        temp_file = os.path.join(tempfile.gettempdir(), filename)
        arquivo.save(temp_file)
        
        # Lê o arquivo Excel
        df = pd.read_excel(temp_file)
        
        # Verifica colunas necessárias
        colunas_necessarias = ['Nome', 'Cargo', 'Departamento', 'Status', 'Data Admissão']
        for coluna in colunas_necessarias:
            if coluna not in df.columns:
                flash(f'Coluna obrigatória ausente: {coluna}', 'danger')
                return redirect(url_for('colaborador.index'))
        
        # Contadores para relatório final
        contador_inseridos = 0
        contador_erros = 0
        erros = []
        
        # Processa cada linha
        for index, row in df.iterrows():
            try:
                # Validar dados básicos
                nome = str(row['Nome']).strip()
                cargo_nome = str(row['Cargo']).strip()
                departamento_nome = str(row['Departamento']).strip()
                status = str(row['Status']).strip()
                
                # Se algum campo obrigatório estiver vazio, pula
                if not nome or nome == 'nan' or not cargo_nome or cargo_nome == 'nan' or not departamento_nome or departamento_nome == 'nan':
                    erros.append(f"Linha {index+2}: Campos obrigatórios faltando")
                    contador_erros += 1
                    continue
                
                # Procura o cargo e departamento pelo nome
                cargo = Cargo.query.filter_by(nome=cargo_nome).first()
                if not cargo:
                    # Tenta criar o cargo se não existir
                    cargo = Cargo(nome=cargo_nome, status='Ativo')
                    cargo.save()
                
                departamento = Departamento.query.filter_by(nome=departamento_nome).first()
                if not departamento:
                    # Tenta criar o departamento se não existir
                    departamento = Departamento(nome=departamento_nome, status='Ativo')
                    departamento.save()
                
                # Tratamento de datas
                try:
                    data_admissao = pd.to_datetime(row['Data Admissão']).date()
                except:
                    erros.append(f"Linha {index+2}: Data de Admissão inválida para {nome}")
                    contador_erros += 1
                    continue
                
                data_nascimento = None
                if 'Data Nascimento' in df.columns and pd.notna(row['Data Nascimento']):
                    try:
                        data_nascimento = pd.to_datetime(row['Data Nascimento']).date()
                    except:
                        pass  # Não obrigatório, ignora erro
                
                # Campos opcionais
                telefone = str(row['Telefone']) if 'Telefone' in df.columns and pd.notna(row['Telefone']) else None
                email = str(row['E-mail']) if 'E-mail' in df.columns and pd.notna(row['E-mail']) else None
                observacoes = str(row['Observações']) if 'Observações' in df.columns and pd.notna(row['Observações']) else None
                
                # Verifica status válido
                status_validos = ['Ativo', 'Inativo', 'Afastado', 'Férias', 'Demitido']
                if status not in status_validos:
                    status = 'Ativo'  # Default
                
                # Verifica CPF duplicado
                cpf = str(row['CPF']) if 'CPF' in df.columns and pd.notna(row['CPF']) else None
                if cpf:
                    colaborador_existente = Colaborador.query.filter_by(cpf=cpf).first()
                    if colaborador_existente:
                        erros.append(f"Linha {index+2}: CPF {cpf} já cadastrado para {colaborador_existente.nome}")
                        contador_erros += 1
                        continue
                
                # Cria o colaborador
                novo_colaborador = Colaborador(
                    nome=nome,
                    cpf=cpf,
                    rg=str(row['RG']) if 'RG' in df.columns and pd.notna(row['RG']) else None,
                    data_nascimento=data_nascimento,
                    data_admissao=data_admissao,
                    cargo_id=cargo.id,
                    departamento_id=departamento.id,
                    status=status,
                    telefone=telefone,
                    email=email,
                    endereco=str(row['Endereço']) if 'Endereço' in df.columns and pd.notna(row['Endereço']) else None,
                    observacoes=observacoes,
                    usuario_id=current_user.id
                )
                
                novo_colaborador.save()
                contador_inseridos += 1
                
            except Exception as e:
                erros.append(f"Linha {index+2}: Erro ao processar - {str(e)}")
                contador_erros += 1
        
        # Remove o arquivo temporário
        os.remove(temp_file)
        
        # Mensagem de conclusão
        if contador_inseridos > 0:
            flash(f'{contador_inseridos} colaboradores importados com sucesso.', 'success')
        
        if contador_erros > 0:
            flash(f'{contador_erros} colaboradores não puderam ser importados.', 'warning')
            for erro in erros[:5]:  # Limita a 5 mensagens para não sobrecarregar
                flash(erro, 'warning')
            if len(erros) > 5:
                flash(f'... e mais {len(erros) - 5} erros.', 'warning')
        
        return redirect(url_for('colaborador.index'))
        
    except Exception as e:
        flash(f'Erro ao processar o arquivo: {str(e)}', 'danger')
        return redirect(url_for('colaborador.index'))

@colaborador_bp.route('/download_template')
@login_required
def download_template():
    """
    Gera e faz download de um template Excel para importação de colaboradores
    """
    # Criar um novo workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Colaboradores"
    
    # Definir cabeçalhos e formatos
    cabecalhos = [
        "Nome", "CPF", "RG", "Data Nascimento", "Data Admissão", 
        "Cargo", "Departamento", "Status", "Telefone", "E-mail", 
        "Endereço", "Observações"
    ]
    
    # Obter cargos e departamentos atuais para referência
    cargos = [cargo.nome for cargo in Cargo.query.filter_by(status='Ativo').all()]
    departamentos = [depto.nome for depto in Departamento.query.filter_by(status='Ativo').all()]
    
    # Adicionar cabeçalhos
    for col_num, header in enumerate(cabecalhos, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # Adicionar exemplos na linha 2
    ws.cell(row=2, column=1, value="João da Silva")
    ws.cell(row=2, column=2, value="123.456.789-00")
    ws.cell(row=2, column=3, value="12.345.678-9")
    ws.cell(row=2, column=4, value=datetime(1980, 1, 1).date())
    ws.cell(row=2, column=5, value=datetime.now().date())
    
    # Usar o primeiro cargo e departamento como exemplo, se disponíveis
    cargo_exemplo = cargos[0] if cargos else "Analista"
    depto_exemplo = departamentos[0] if departamentos else "Recursos Humanos"
    
    ws.cell(row=2, column=6, value=cargo_exemplo)
    ws.cell(row=2, column=7, value=depto_exemplo)
    ws.cell(row=2, column=8, value="Ativo")
    ws.cell(row=2, column=9, value="(11) 98765-4321")
    ws.cell(row=2, column=10, value="joao.silva@exemplo.com")
    ws.cell(row=2, column=11, value="Rua Exemplo, 123 - São Paulo/SP")
    ws.cell(row=2, column=12, value="Colaborador excelente")
    
    # Adicionar instruções na linha 3 (em itálico)
    for col_num, header in enumerate(cabecalhos, 1):
        if header in ["Nome", "Data Admissão", "Cargo", "Departamento", "Status"]:
            cell = ws.cell(row=3, column=col_num, value="(Obrigatório)")
        else:
            cell = ws.cell(row=3, column=col_num, value="(Opcional)")
        cell.font = Font(italic=True)
        cell.alignment = Alignment(horizontal='center')
    
    # Ajustar largura das colunas
    for col_num, _ in enumerate(cabecalhos, 1):
        column_letter = openpyxl.utils.get_column_letter(col_num)
        ws.column_dimensions[column_letter].width = 20
    
    # Criar uma nova planilha para referência de cargos e departamentos
    ws_ref = wb.create_sheet(title="Referências")
    
    # Adicionar lista de cargos
    ws_ref.cell(row=1, column=1, value="Cargos Disponíveis")
    ws_ref.cell(row=1, column=1).font = Font(bold=True)
    
    for i, cargo in enumerate(cargos, 2):
        ws_ref.cell(row=i, column=1, value=cargo)
    
    # Adicionar lista de departamentos
    ws_ref.cell(row=1, column=3, value="Departamentos Disponíveis")
    ws_ref.cell(row=1, column=3).font = Font(bold=True)
    
    for i, depto in enumerate(departamentos, 2):
        ws_ref.cell(row=i, column=3, value=depto)
    
    # Adicionar lista de status
    ws_ref.cell(row=1, column=5, value="Status Disponíveis")
    ws_ref.cell(row=1, column=5).font = Font(bold=True)
    
    status_list = ["Ativo", "Inativo", "Afastado", "Férias", "Demitido"]
    for i, status in enumerate(status_list, 2):
        ws_ref.cell(row=i, column=5, value=status)
    
    # Ajustar largura das colunas na planilha de referências
    ws_ref.column_dimensions['A'].width = 25
    ws_ref.column_dimensions['C'].width = 25
    ws_ref.column_dimensions['E'].width = 25
    
    # Salvar para um stream de bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        download_name='template_colaboradores.xlsx',
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@colaborador_bp.route('/api/buscar')
@login_required
def api_buscar():
    """
    API para buscar colaboradores (para uso em selects dinâmicos)
    """
    termo = request.args.get('termo', '')
    status = request.args.get('status', 'Ativo')
    
    query = Colaborador.query
    
    if status:
        query = query.filter_by(status=status)
    
    if termo:
        query = query.filter(Colaborador.nome.ilike(f'%{termo}%'))
    
    colaboradores = query.limit(10).all()
    
    resultado = [
        {
            'id': c.id,
            'nome': c.nome,
            'cargo': c.cargo.nome if c.cargo else '',
            'departamento': c.departamento.nome if c.departamento else ''
        }
        for c in colaboradores
    ]
    
    return jsonify(resultado)

@colaborador_bp.route('/<int:id>/dados_bancarios', methods=['GET'])
@login_required
def dados_bancarios(id):
    """
    Retorna os dados bancários de um colaborador (JSON)
    """
    colaborador = Colaborador.query.get_or_404(id)
    dados = colaborador.dados_bancarios
    
    if dados:
        return jsonify({
            'id': dados.id,
            'pix': dados.pix or '',
            'banco': dados.banco or '',
            'agencia': dados.agencia or '',
            'conta': dados.conta or ''
        })
    else:
        return jsonify({
            'id': None,
            'pix': '',
            'banco': '',
            'agencia': '',
            'conta': ''
        })

@colaborador_bp.route('/<int:id>/dados_bancarios', methods=['POST'])
@login_required
def salvar_dados_bancarios(id):
    """
    Salva ou atualiza os dados bancários de um colaborador
    """
    colaborador = Colaborador.query.get_or_404(id)
    
    pix = request.form.get('pix', '').strip()
    banco = request.form.get('banco', '').strip()
    agencia = request.form.get('agencia', '').strip()
    conta = request.form.get('conta', '').strip()
    
    # Validação: pelo menos um campo deve ser preenchido
    if not pix and not banco and not agencia and not conta:
        flash('Por favor, preencha pelo menos um campo de dados bancários.', 'danger')
        return redirect(url_for('colaborador.index'))
    
    # Permite valores None para campos não preenchidos
    pix = pix if pix else None
    banco = banco if banco else None
    agencia = agencia if agencia else None
    conta = conta if conta else None
    
    # Verifica se já existem dados bancários
    dados_existentes = colaborador.dados_bancarios
    
    if dados_existentes:
        # Atualiza dados existentes
        dados_existentes.pix = pix
        dados_existentes.banco = banco
        dados_existentes.agencia = agencia
        dados_existentes.conta = conta
        dados_existentes.atualizado_em = datetime.now()
        db.session.commit()
        flash('Dados bancários atualizados com sucesso.', 'success')
    else:
        # Cria novos dados bancários
        novos_dados = DadosBancarios(
            colaborador_id=colaborador.id,
            pix=pix,
            banco=banco,
            agencia=agencia,
            conta=conta
        )
        db.session.add(novos_dados)
        db.session.commit()
        flash('Dados bancários cadastrados com sucesso.', 'success')
    
    return redirect(url_for('colaborador.index'))


@colaborador_bp.route('/<int:id>/mudanca_funcao', methods=['POST'])
@login_required
def salvar_mudanca_funcao(id):
    """
    Registra mudanças em ``mudanca_funcao`` (lista ``{data, funcao_id}``), migrando legado
    ``dados_a``. A partir de cada data vale a função indicada; antes da primeira data
    continua valendo ``colaboradores.cargo_id`` (não alterado por este fluxo).
    """
    colaborador = Colaborador.query.get_or_404(id)

    def _parse_json_lista(campo):
        raw = (request.form.get(campo) or '').strip()
        if not raw:
            return []
        try:
            parsed = json.loads(raw)
            return parsed if isinstance(parsed, list) else []
        except (json.JSONDecodeError, TypeError):
            return []

    def _resolver_nome_funcao_para_id(nome: str):
        c = Cargo.query.filter_by(nome=(nome or '').strip()).first()
        return c.id if c else None

    def _item_historico_bate_exclusao(item, exc):
        if not isinstance(exc, dict):
            return False
        data_item = (item.get('data') or '').strip()
        data_exc = (exc.get('data') or '').strip()
        if data_item != data_exc:
            return False
        id_exc = normalizar_id_opcional(exc.get('funcao_id'))
        id_item = normalizar_id_opcional(item.get('funcao_id'))
        if id_exc is not None:
            return id_item == id_exc
        nome_exc = (exc.get('funcao') or '').strip()
        if nome_exc:
            c = Cargo.query.filter_by(nome=nome_exc).first()
            return c is not None and id_item == c.id
        return False

    exclusoes = _parse_json_lista('mudancas_excluir_json')

    mudancas_json = (request.form.get('mudancas_funcao_json') or '').strip()
    mudancas_recebidas = []

    if mudancas_json:
        try:
            parsed = json.loads(mudancas_json)
            if isinstance(parsed, list):
                mudancas_recebidas = parsed
        except (json.JSONDecodeError, TypeError):
            mudancas_recebidas = []

    mudancas_para_salvar = []
    for item in mudancas_recebidas:
        if not isinstance(item, dict):
            continue
        data_mudanca_str = (item.get('data') or '').strip()
        cargo_id = item.get('cargo_id')
        if not data_mudanca_str or not cargo_id:
            continue
        try:
            cargo_id = int(cargo_id)
        except (TypeError, ValueError):
            flash('Uma das funções selecionadas é inválida.', 'danger')
            return redirect(url_for('colaborador.index'))
        try:
            data_mudanca = datetime.strptime(data_mudanca_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Uma das datas informadas é inválida.', 'danger')
            return redirect(url_for('colaborador.index'))

        cargo = Cargo.query.get(cargo_id)
        if not cargo:
            flash('Uma das funções selecionadas é inválida.', 'danger')
            return redirect(url_for('colaborador.index'))

        mudancas_para_salvar.append({
            'data': data_mudanca.strftime('%Y-%m-%d'),
            'cargo_id': cargo.id
        })

    # Fallback para formulário antigo (uma mudança por envio)
    if not mudancas_para_salvar and not exclusoes:
        data_mudanca_str = (request.form.get('data_mudanca') or '').strip()
        cargo_id = request.form.get('cargo_id')
        if not data_mudanca_str or not cargo_id:
            flash('Informe data e função para registrar a mudança.', 'danger')
            return redirect(url_for('colaborador.index'))
        try:
            cargo_id = int(cargo_id)
        except (TypeError, ValueError):
            flash('Função selecionada é inválida.', 'danger')
            return redirect(url_for('colaborador.index'))
        try:
            data_mudanca = datetime.strptime(data_mudanca_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Data de mudança de função inválida.', 'danger')
            return redirect(url_for('colaborador.index'))
        cargo = Cargo.query.get(cargo_id)
        if not cargo:
            flash('Função selecionada é inválida.', 'danger')
            return redirect(url_for('colaborador.index'))
        mudancas_para_salvar.append({
            'data': data_mudanca.strftime('%Y-%m-%d'),
            'cargo_id': cargo.id
        })

    dados_adicionais = colaborador.get_dados_adicionais_dict()
    historico = historico_mudanca_funcao_normalizado(
        dados_adicionais, resolver_nome_funcao=_resolver_nome_funcao_para_id
    )

    for exc in exclusoes:
        historico = [h for h in historico if not _item_historico_bate_exclusao(h, exc)]

    for mudanca in mudancas_para_salvar:
        historico.append({
            'data': mudanca['data'],
            'funcao_id': mudanca['cargo_id']
        })

    historico.sort(key=lambda item: item.get('data', ''), reverse=True)

    dados_adicionais.pop('dados_a', None)
    dados_adicionais['mudanca_funcao'] = historico
    colaborador.dados_adicionais = json.dumps(dados_adicionais, ensure_ascii=False)
    # cargo_id permanece a função “base” (vigente antes da primeira data no histórico);
    # a partir das datas em mudanca_funcao usa-se funcao_id_vigente_em / get_funcao().
    colaborador.atualizado_em = datetime.now()
    colaborador.save()

    if mudancas_para_salvar and exclusoes:
        flash('Mudanças de função atualizadas com sucesso.', 'success')
    elif mudancas_para_salvar:
        if len(mudancas_para_salvar) == 1:
            flash('Mudança de função registrada com sucesso.', 'success')
        else:
            flash(f'{len(mudancas_para_salvar)} mudanças de função registradas com sucesso.', 'success')
    elif exclusoes:
        flash('Mudanças de função removidas com sucesso.', 'success')
    return redirect(url_for('colaborador.index'))

