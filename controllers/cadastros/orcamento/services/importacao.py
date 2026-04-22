"""Service de importação de orçamento via planilha."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy import func
from werkzeug.datastructures import FileStorage

from models.database import db
from models.orcamento import ItemOrcamento, ItemOrcamentoReferenciaMaterial, Orcamento


def _normalizar_data(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _to_decimal(valor, default=Decimal("0.00")):
    if valor in (None, ""):
        return default
    try:
        return Decimal(str(valor))
    except Exception:
        return default


def _to_int(valor, default=1):
    if valor in (None, ""):
        return default
    try:
        return int(float(valor))
    except Exception:
        return default


def _buscar_material_por_texto_item(texto_item: str):
    if not texto_item:
        return None
    texto_item = str(texto_item).strip()
    if not texto_item:
        return None
    ref = (
        ItemOrcamentoReferenciaMaterial.query.filter(
            func.lower(ItemOrcamentoReferenciaMaterial.texto_item) == texto_item.lower()
        )
        .order_by(ItemOrcamentoReferenciaMaterial.id.desc())
        .first()
    )
    if ref:
        return ref.material_id

    db.session.add(
        ItemOrcamentoReferenciaMaterial(
            texto_item=texto_item,
            material_id=None,
        )
    )
    return None


def importar_planilha_orcamento(
    arquivo: FileStorage,
    usuario_id: int | None = None,
    orcamento_id_existente: int | None = None,
):
    if not arquivo.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise ValueError("Formato inválido. Envie um arquivo Excel (.xlsx).")

    stream = BytesIO(arquivo.read())
    wb = load_workbook(filename=stream, data_only=True, read_only=True)

    if "ENTRADAS" not in wb.sheetnames:
        raise ValueError("Aba 'ENTRADAS' não encontrada.")
    if "INSUMOS" not in wb.sheetnames:
        raise ValueError("Aba 'INSUMOS' não encontrada.")

    ws_entradas = wb["ENTRADAS"]
    ws_insumos = wb["INSUMOS"]

    if orcamento_id_existente is not None:
        orcamento = Orcamento.query.get(orcamento_id_existente)
        if not orcamento:
            raise ValueError(f"Orçamento {orcamento_id_existente} não encontrado.")
    else:
        nome = (ws_entradas["F4"].value or "").strip()
        if not nome:
            raise ValueError("Nome do orçamento não encontrado em ENTRADAS!F4.")
        data_orc = _normalizar_data(ws_entradas["O1"].value)

        orcamento = Orcamento(
            nome=nome,
            data=data_orc,
            status="Importado",
            usuario_id=usuario_id,
        )
        db.session.add(orcamento)
        db.session.flush()

    col_c = column_index_from_string("C")
    col_d = column_index_from_string("D")
    col_k = column_index_from_string("K")
    col_o = column_index_from_string("O")
    col_q = column_index_from_string("Q")

    itens_importados = 0
    grupo_atual = None
    for linha in range(2, (ws_insumos.max_row or 1) + 1):
        tipo_linha_raw = ws_insumos.cell(row=linha, column=col_c).value
        tipo_linha = str(tipo_linha_raw).strip() if tipo_linha_raw not in (None, "") else ""

        descricao_raw = ws_insumos.cell(row=linha, column=col_d).value
        if tipo_linha == "ítem":
            grupo_texto = str(descricao_raw).strip() if descricao_raw not in (None, "") else ""
            grupo_atual = grupo_texto or None
            continue

        k_raw = ws_insumos.cell(row=linha, column=col_k).value
        if descricao_raw in (None, "") or k_raw in (None, ""):
            continue

        descricao_item = str(descricao_raw).strip()
        if not descricao_item:
            continue

        quantidade_raw = ws_insumos.cell(row=linha, column=col_k).value
        valor_raw = ws_insumos.cell(row=linha, column=col_q).value

        item = ItemOrcamento(
            orcamento_id=orcamento.id,
            descricao_item=descricao_item,
            grupo=grupo_atual,
            material_id=_buscar_material_por_texto_item(descricao_item),
            quantidade=max(1, _to_int(quantidade_raw, default=1)),
            valor=_to_decimal(valor_raw, default=Decimal("0.00")),
        )
        db.session.add(item)
        itens_importados += 1

    if itens_importados == 0:
        raise ValueError("Nenhum item válido foi encontrado na aba INSUMOS.")

    db.session.commit()
    return {
        "orcamento_id": orcamento.id,
        "nome": orcamento.nome,
        "itens_importados": itens_importados,
        "vinculado_a_existente": orcamento_id_existente is not None,
    }
