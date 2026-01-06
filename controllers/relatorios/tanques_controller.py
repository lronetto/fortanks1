from flask import Blueprint, render_template, request, send_file, jsonify, flash, redirect, url_for
from flask_login import login_required
from controllers.relatorios.script_email import enviar_relatorio_por_email
from models.tanque import TanquesPecas, Tanques
from models.contrato import Contrato
from models.database import db
from datetime import datetime
import pandas as pd
import io
import json
from weasyprint import HTML, CSS
import os
from sqlalchemy import func, outerjoin
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, NamedStyle
from openpyxl.styles.numbers import FORMAT_PERCENTAGE_00
from controllers.relatorios.script_email import relatorio_semanal
tanques_bp = Blueprint('relatorios_tanques', __name__, url_prefix='/relatorios/tanques')

def get_dados_tanques(contrato_id=None, tanque_id=None):
    """
    Busca dados de tanques com filtros opcionais
    """
    query = Tanques.query
    
    # Aplicar filtros
    if tanque_id:
        query = query.filter(Tanques.id == tanque_id)
    
    if contrato_id:
        query = query.filter(Tanques.contrato_id == contrato_id)
    
    tanques = query.order_by(Tanques.nome).all()
    
    dados = []
    for tanque in tanques:
        # Calcular quantidade prevista: (placas_normais + placas_fecho) * quantidade
        placas_normais = tanque.placas_normais or 0
        placas_fecho = tanque.placas_fecho or 0
        quantidade_tanques = tanque.quantidade or 1
        quantidade_prevista = (placas_normais + placas_fecho) * quantidade_tanques
        
        # Buscar peças concretadas (com data_concretagem não nula)
        pecas_concretadas = TanquesPecas.query.filter(
            TanquesPecas.tanque_id == tanque.id,
            TanquesPecas.data_concretagem.isnot(None),
        ).count()
        
        # Buscar peças com acabamento
        pecas_acabadas = 0
        for peca in tanque.TanquesPecas:
            if peca.qualidade:
                try:
                    qualidade_dict = json.loads(peca.qualidade) if isinstance(peca.qualidade, str) else peca.qualidade
                    if 'acabamento' in qualidade_dict and qualidade_dict['acabamento']:
                        pecas_acabadas += 1
                except:
                    pass
        
        # Calcular percentual de conclusão
        percentual_conclusao = (pecas_concretadas / quantidade_prevista * 100) if quantidade_prevista > 0 else 0
        
        dados.append({
            'id': tanque.id,
            'nome': tanque.nome,
            'contrato': tanque.contrato.nome if tanque.contrato else 'Sem contrato',
            'contrato_id': tanque.contrato_id,
            'sistema': tanque.sistema,
            'dimensoes': tanque.dimensoes,
            'altura_total': tanque.altura_total,
            'altura_util': tanque.altura_util,
            'quantidade': quantidade_tanques,
            'placas_normais': placas_normais,
            'placas_fecho': placas_fecho,
            'quantidade_prevista': quantidade_prevista,
            'pecas_concretadas': pecas_concretadas,
            'pecas_acabadas': pecas_acabadas,
            'percentual_conclusao': round(percentual_conclusao, 2),
            'bainhas': tanque.quantidade_bainhas or 0,
            'item_nf': tanque.item_nf or '-',
        })
    
    return dados

def agrupar_por_projeto(dados):
    """
    Agrupa dados de tanques por projeto
    """
    projetos = {}
    
    for item in dados:
        projeto_nome = item['contrato']
        if projeto_nome not in projetos:
            projetos[projeto_nome] = {
                'quantidade_tanques': 0,
                'quantidade_prevista': 0,
                'pecas_concretadas': 0,
                'pecas_acabadas': 0,
            }
        
        projetos[projeto_nome]['quantidade_tanques'] += item['quantidade']
        projetos[projeto_nome]['quantidade_prevista'] += item['quantidade_prevista']
        projetos[projeto_nome]['pecas_concretadas'] += item['pecas_concretadas']
        projetos[projeto_nome]['pecas_acabadas'] += item['pecas_acabadas']
    
    # Calcular percentuais
    for projeto_nome, dados_projeto in projetos.items():
        if dados_projeto['quantidade_prevista'] > 0:
            dados_projeto['percentual_conclusao'] = round(
                (dados_projeto['pecas_concretadas'] / dados_projeto['quantidade_prevista']) * 100, 2
            )
        else:
            dados_projeto['percentual_conclusao'] = 0
    
    return projetos

@tanques_bp.route('/teste1', methods=['GET'])
@login_required
def teste1():
    """Teste 1"""
    relatorio_semanal()
    return 'Teste 1'

@tanques_bp.route('/', methods=['GET'])
@login_required
def index():
    """Página principal do relatório de tanques"""
    contrato_id = request.args.get('contrato_id', type=int)
    
    # Buscar contratos ativos
    contratos = Contrato.query.filter_by(ativo=True).order_by(Contrato.nome).all()
    
    # Buscar tanques para filtro
    tanques_query = Tanques.query.outerjoin(Contrato)
    if contrato_id:
        tanques_query = tanques_query.filter(Tanques.contrato_id == contrato_id)
    tanques = tanques_query.order_by(Tanques.nome).all()
    
    # Buscar dados
    dados = get_dados_tanques(contrato_id=contrato_id)
    projetos = agrupar_por_projeto(dados)
    
    # Preparar dados para gráficos
    labels_projetos = list(projetos.keys()) if projetos else []
    valores_conclusao = [projetos[p]['percentual_conclusao'] for p in labels_projetos] if labels_projetos else []
    valores_previsto = [projetos[p]['quantidade_prevista'] for p in labels_projetos] if labels_projetos else []
    valores_realizado = [projetos[p]['pecas_concretadas'] for p in labels_projetos] if labels_projetos else []
    
    # Calcular totais
    total_tanques = sum([item['quantidade'] for item in dados])
    total_previsto = sum([item['quantidade_prevista'] for item in dados])
    total_concretadas = sum([item['pecas_concretadas'] for item in dados])
    total_acabadas = sum([item['pecas_acabadas'] for item in dados])
    percentual_geral = round((total_concretadas / total_previsto * 100) if total_previsto > 0 else 0, 2)
    
    return render_template(
        'relatorios/tanques/index.html',
        dados=dados,
        projetos=projetos,
        contratos=contratos,
        tanques=tanques,
        contrato_id=contrato_id,
        labels_projetos=labels_projetos,
        valores_conclusao=valores_conclusao,
        valores_previsto=valores_previsto,
        valores_realizado=valores_realizado,
        total_tanques=total_tanques,
        total_previsto=total_previsto,
        total_concretadas=total_concretadas,
        total_acabadas=total_acabadas,
        percentual_geral=percentual_geral
    )

@tanques_bp.route('/api/tanques-por-projeto', methods=['GET'])
@login_required
def api_tanques_por_projeto():
    """API para buscar tanques filtrados por projeto"""
    contrato_id = request.args.get('contrato_id', type=int)
    
    query = Tanques.query.outerjoin(Contrato)
    
    if contrato_id:
        query = query.filter(Tanques.contrato_id == contrato_id)
    
    tanques = query.order_by(Tanques.nome).all()
    
    tanques_json = [{'id': t.id, 'nome': t.nome} for t in tanques]
    
    return jsonify({'tanques': tanques_json})

@tanques_bp.route('/api/dados', methods=['GET'])
@login_required
def api_dados():
    """Endpoint AJAX para buscar dados do relatório"""
    contrato_id = request.args.get('contrato_id', type=int)
    tanque_id = request.args.get('tanque_id', type=int)
    
    dados = get_dados_tanques(contrato_id=contrato_id, tanque_id=tanque_id)
    projetos = agrupar_por_projeto(dados)
    
    # Calcular totais
    total_tanques = sum([item['quantidade'] for item in dados])
    total_previsto = sum([item['quantidade_prevista'] for item in dados])
    total_concretadas = sum([item['pecas_concretadas'] for item in dados])
    total_acabadas = sum([item['pecas_acabadas'] for item in dados])
    percentual_geral = round((total_concretadas / total_previsto * 100) if total_previsto > 0 else 0, 2)
    
    # Preparar dados para gráficos
    labels_projetos = list(projetos.keys()) if projetos else []
    valores_conclusao = [projetos[p]['percentual_conclusao'] for p in labels_projetos] if labels_projetos else []
    valores_previsto = [projetos[p]['quantidade_prevista'] for p in labels_projetos] if labels_projetos else []
    valores_realizado = [projetos[p]['pecas_concretadas'] for p in labels_projetos] if labels_projetos else []
    
    return jsonify({
        'dados': dados,
        'projetos': projetos,
        'labels_projetos': labels_projetos,
        'valores_conclusao': valores_conclusao,
        'valores_previsto': valores_previsto,
        'valores_realizado': valores_realizado,
        'total_tanques': total_tanques,
        'total_previsto': total_previsto,
        'total_concretadas': total_concretadas,
        'total_acabadas': total_acabadas,
        'percentual_geral': percentual_geral
    })

@tanques_bp.route('/exportar/excel', methods=['GET'])
@login_required
def exportar_excel():
    """Exportar relatório em formato Excel"""
    contrato_id = request.args.get('contrato_id', type=int)
    tanque_id = request.args.get('tanque_id', type=int)
    
    dados = get_dados_tanques(contrato_id=contrato_id, tanque_id=tanque_id)
    
    if not dados:
        from flask import flash, redirect, url_for
        flash('Nenhum dado encontrado para exportar.', 'warning')
        return redirect(url_for('relatorios_tanques.index'))
    
    # Criar DataFrame
    df = pd.DataFrame(dados)
    
    # Renomear colunas
    df.rename(columns={
        'nome': 'Tanque',
        'contrato': 'Projeto',
        'sistema': 'Sistema',
        'dimensoes': 'Dimensões',
        'altura_total': 'Altura Total (m)',
        'altura_util': 'Altura Útil (m)',
        'quantidade': 'Quantidade',
        'placas_normais': 'Placas Normais',
        'placas_fecho': 'Placas Fecho',
        'quantidade_prevista': 'Quantidade Prevista',
        'pecas_concretadas': 'Pecas Concretadas',
        'pecas_acabadas': 'Pecas Acabadas',
        'percentual_conclusao': '% Conclusão',
        'bainhas': 'Bainhas',
        'item_nf': 'Item NF',
    }, inplace=True)
    
    # Selecionar colunas para exportação
    colunas_exportacao = [
        'Tanque', 'Projeto', 'Sistema', 'Dimensões', 'Altura Total (m)', 'Altura Útil (m)',
        'Quantidade', 'Placas Normais', 'Placas Fecho', 'Quantidade Prevista',
        'Pecas Concretadas', 'Pecas Acabadas', '% Conclusão', 'Bainhas', 'Item NF'
    ]
    df_export = df[colunas_exportacao]
    
    # Adicionar linha de totais
    totais = {
        'Tanque': 'TOTAL',
        'Projeto': '',
        'Sistema': '',
        'Dimensões': '',
        'Altura Total (m)': '',
        'Altura Útil (m)': '',
        'Quantidade': df['Quantidade'].sum(),
        'Placas Normais': df['Placas Normais'].sum(),
        'Placas Fecho': df['Placas Fecho'].sum(),
        'Quantidade Prevista': df['Quantidade Prevista'].sum(),
        'Pecas Concretadas': df['Pecas Concretadas'].sum(),
        'Pecas Acabadas': df['Pecas Acabadas'].sum(),
        '% Conclusão': round((df['Pecas Concretadas'].sum() / df['Quantidade Prevista'].sum() * 100) if df['Quantidade Prevista'].sum() > 0 else 0, 2),
        'Bainhas': df['Bainhas'].sum(),
        'Item NF': '',
    }
    
    df_totais = pd.DataFrame([totais])
    df_export = pd.concat([df_export, df_totais], ignore_index=True)
    
    # Criar Excel na memória
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
        df_export.to_excel(writer, index=False, sheet_name='Relatório de Tanques')
        
        # Ajustar largura das colunas
        worksheet = writer.sheets['Relatório de Tanques']
        for idx, col in enumerate(df_export.columns):
            max_length = max(
                df_export[col].astype(str).apply(len).max(),
                len(col)
            )
            adjusted_width = min(max_length + 2, 50)
            worksheet.set_column(idx, idx, adjusted_width)
        
        # Formatar linha de totais
        last_row = len(df_export)
        header_format = writer.book.add_format({'bold': True, 'bg_color': '#D3D3D3'})
        for col_idx in range(len(df_export.columns)):
            worksheet.write(last_row, col_idx, df_export.iloc[last_row - 1, col_idx], header_format)
    
    output.seek(0)
    
    # Nome do arquivo
    nome_arquivo = f"relatorio_tanques_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nome_arquivo
    )

@tanques_bp.route('/relatorio-projeto-excel/<int:contrato_id>', methods=['GET'])
@login_required
def relatorio_projeto_excel(contrato_id):
    """
    Gera relatório em Excel com tanques do projeto, quantidade prevista (PN+PF) e realizada (concretadas)
    """
    try:
        # Buscar o contrato
        contrato = Contrato.query.get_or_404(contrato_id)
        
        # Buscar todos os tanques do contrato
        tanques = Tanques.query.filter_by(contrato_id=contrato_id).order_by(Tanques.nome).all()
        
        if not tanques:
            flash('Nenhum tanque encontrado para este projeto.', 'warning')
            return redirect(url_for('tanque.index'))
        
        # Preparar dados para o relatório
        dados_relatorio = []
        
        for tanque in tanques:
            # Calcular quantidade prevista: (placas_normais + placas_fecho) * quantidade
            placas_normais = tanque.placas_normais or 0
            placas_fecho = tanque.placas_fecho or 0
            quantidade_tanques = tanque.quantidade or 1
            quantidade_prevista = (placas_normais + placas_fecho) * quantidade_tanques
            
            # Buscar peças concretadas (com data_concretagem não nula)
            pecas_concretadas = TanquesPecas.query.filter(
                TanquesPecas.tanque_id == tanque.id,
                TanquesPecas.data_concretagem.isnot(None),
            ).count()
            
            percentual_concluido = (pecas_concretadas/quantidade_prevista)*100 if quantidade_prevista > 0 else 0
            dados_relatorio.append({
                'ID': tanque.id,
                'Tanque': tanque.nome,
                'Placas': quantidade_prevista,
                'Quantidade Realizada (Concretadas)': pecas_concretadas,
                'Concluido': percentual_concluido / 100,  # Dividir por 100 para formato de porcentagem (0.85 = 85%)
            })
        
        # Criar DataFrame
        df = pd.DataFrame(dados_relatorio)
        
        if df.empty:
            flash('Nenhum dado encontrado para este projeto.', 'warning')
            return redirect(url_for('tanque.index'))
        
        # Adicionar linha de totais
        percentual_total = (df['Quantidade Realizada (Concretadas)'].sum()/df['Placas'].sum()*100) if df['Placas'].sum() > 0 else 0
        totais = {
            'ID': '',
            'Tanque': 'TOTAL',
            'Placas': df['Placas'].sum(),
            'Quantidade Realizada (Concretadas)': df['Quantidade Realizada (Concretadas)'].sum(),
            'Concluido': percentual_total / 100  # Dividir por 100 para formato de porcentagem
        }
        
        # Adicionar linha de totais ao DataFrame
        df_totais = pd.DataFrame([totais])
        df = pd.concat([df, df_totais], ignore_index=True)
        
        # Criar Excel na memória
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Relatório de Tanques')
            
            # Ajustar largura das colunas
            worksheet = writer.sheets['Relatório de Tanques']
            for idx, col in enumerate(df.columns):
                max_length = max(
                    df[col].astype(str).apply(len).max(),
                    len(col)
                )
                # Limitar largura máxima
                adjusted_width = min(max_length + 2, 50)
                col_letter = get_column_letter(idx + 1)
                worksheet.column_dimensions[col_letter].width = adjusted_width
            
            # Formatar coluna de "Concluido" como porcentagem
            col_concluido_idx = list(df.columns).index('Concluido')
            col_concluido_letter = get_column_letter(col_concluido_idx + 1)
            
            # Aplicar formatação de porcentagem em todas as linhas (exceto cabeçalho)
            for row_idx in range(2, len(df) + 2):  # Começa na linha 2 (após cabeçalho)
                cell = worksheet[f'{col_concluido_letter}{row_idx}']
                cell.number_format = '0.00%'  # Formato: 85.50%
            
            # Formatar linha de totais (última linha)
            last_row = len(df) + 1  # +1 porque o Excel começa em 1 e tem cabeçalho
            for col_idx, col in enumerate(df.columns):
                col_letter = get_column_letter(col_idx + 1)
                cell = worksheet[f'{col_letter}{last_row}']
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
                # Se for a coluna "Concluido", aplicar formatação de porcentagem também
                if col == 'Concluido':
                    cell.number_format = '0.00%'
        
        output.seek(0)
        
        # Nome do arquivo
        nome_arquivo = f"relatorio_tanques_{contrato.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=nome_arquivo
        )
        
    except Exception as e:
        flash(f'Erro ao gerar relatório: {str(e)}', 'danger')
        return redirect(url_for('tanque.index'))

@tanques_bp.route('/exportar/pdf', methods=['GET'])
@login_required
def exportar_pdf():
    """Exportar relatório em formato PDF"""
    contrato_id = request.args.get('contrato_id', type=int)
    tanque_id = request.args.get('tanque_id', type=int)
    
    dados = get_dados_tanques(contrato_id=contrato_id, tanque_id=tanque_id)
    projetos = agrupar_por_projeto(dados)
    
    # Calcular totais
    total_tanques = sum([item['quantidade'] for item in dados])
    total_previsto = sum([item['quantidade_prevista'] for item in dados])
    total_concretadas = sum([item['pecas_concretadas'] for item in dados])
    total_acabadas = sum([item['pecas_acabadas'] for item in dados])
    percentual_geral = round((total_concretadas / total_previsto * 100) if total_previsto > 0 else 0, 2)
    
    # Buscar nome do projeto se filtrado
    projeto_nome = None
    if contrato_id:
        contrato = Contrato.query.get(contrato_id)
        if contrato:
            projeto_nome = contrato.nome
    
    # Caminho absoluto da logo para o WeasyPrint
    logo_path = os.path.abspath(os.path.join('static', 'img', 'logo.png'))
    logo_path_uri = 'file:///' + logo_path.replace('\\', '/').replace('\\', '/')
    
    html = render_template(
        'relatorios/tanques/pdf.html',
        dados=dados,
        projetos=projetos,
        projeto_nome=projeto_nome,
        total_tanques=total_tanques,
        total_previsto=total_previsto,
        total_concretadas=total_concretadas,
        total_acabadas=total_acabadas,
        percentual_geral=percentual_geral,
        now=datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        logo_path=logo_path_uri
    )
    
    pdf_bytes = HTML(string=html).write_pdf(
        stylesheets=[CSS(string='body { font-family: Arial, sans-serif;}')]
    )
    
    pdf_io = io.BytesIO(pdf_bytes)
    pdf_io.seek(0)
    
    return send_file(
        pdf_io,
        mimetype='application/pdf',
        as_attachment=False,
        download_name='relatorio_tanques.pdf'
    )

