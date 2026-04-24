"""Service de importação de orçamento via planilha."""

from __future__ import annotations

import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy import func
from werkzeug.datastructures import FileStorage

from models.database import db
from models.orcamento import ItemOrcamento, ItemOrcamentoReferenciaMaterial, Orcamento
from utils.parser import (
    ToInt,
    StrToDate,
    ToDecimal,
)

def _buscar_material_por_texto_item(texto_item: str):
    if not texto_item:
        return None
    texto_item = str(texto_item).strip()
    if not texto_item:
        return None
    ref = (
        ItemOrcamentoReferenciaMaterial.query.filter(
            func.lower(ItemOrcamentoReferenciaMaterial.texto_item) == texto_item.lower()
        )
        .order_by(ItemOrcamentoReferenciaMaterial.id.desc())
        .first()
    )
    if ref:
        return ref.material_id

    db.session.add(
        ItemOrcamentoReferenciaMaterial(
            texto_item=texto_item,
            material_id=None,
        )
    )
    return None


def importar_planilha_orcamento(
    arquivo: FileStorage,
    usuario_id: int | None = None,
    orcamento_id_existente: int | None = None,
):
    if not arquivo.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise ValueError("Formato inválido. Envie um arquivo Excel (.xlsx).")

    stream = BytesIO(arquivo.read())
    wb = load_workbook(filename=stream, data_only=True, read_only=True)

    if "ENTRADAS" not in wb.sheetnames:
        raise ValueError("Aba 'ENTRADAS' não encontrada.")
    if "INSUMOS" not in wb.sheetnames:
        raise ValueError("Aba 'INSUMOS' não encontrada.")

    ws_entradas = wb["ENTRADAS"]
    ws_insumos = wb["INSUMOS"]

    if orcamento_id_existente is not None:
        orcamento = Orcamento.query.get(orcamento_id_existente)
        if not orcamento:
            raise ValueError(f"Orçamento {orcamento_id_existente} não encontrado.")
    else:
        nome = (ws_entradas["F4"].value or "").strip()
        if not nome:
            raise ValueError("Nome do orçamento não encontrado em ENTRADAS!F4.")
        data_orc = ws_entradas["O1"].value or None
        if ws_entradas["C59"].value in ["x","X"]:
            tipo = "SC-10"
        elif ws_entradas["E59"].value in ["x","X"]:
            tipo = "SC-14"
        elif ws_entradas["C64"].value in ["x","X"]:
            tipo = "SC-6"
        dados_adicionais = {
            "Ht": ws_entradas["E15"].value,
            "Hu": ws_entradas["E16"].value,
            "Quantidade": ws_entradas["F21"].value,
            "DProjeto": ws_entradas["G24"].value,
            "DCalculo": ws_entradas["G25"].value,
            "tipo": tipo,
            "Paineis":{
                "PN": ToInt(ws_entradas["G96"].value,0),
                "PF": ToInt(ws_entradas["G90"].value,0) + ToInt(ws_entradas["G92"].value,0),
            }
        }
        dados_adicionais = json.dumps(
            dados_adicionais,
            ensure_ascii=False,
            default=lambda o: o.isoformat() if isinstance(o, (datetime, date)) else str(o),
        )
        orcamento = Orcamento(
            nome=nome,
            data=data_orc,
            status="Importado",
            usuario_id=usuario_id,
            dados_adicionais=dados_adicionais,
        )
        db.session.add(orcamento)
        db.session.flush()

    col_c = column_index_from_string("C")
    col_d = column_index_from_string("D")
    col_k = column_index_from_string("K")
    col_o = column_index_from_string("O")
    col_q = column_index_from_string("Q")
    col_j = column_index_from_string("J")
    col_l = column_index_from_string("L")
    col_m = column_index_from_string("M")
    col_n = column_index_from_string("N")
    itens_importados = 0
    grupo_atual = None
    ultima_linha_insumos = min(ws_insumos.max_row or 1, 277)
    for linha in range(36, ultima_linha_insumos + 1):
        tipo_linha_raw = ws_insumos.cell(row=linha, column=col_c).value
        tipo_linha = str(tipo_linha_raw).strip() if tipo_linha_raw not in (None, "") else ""
        k_raw = ws_insumos.cell(row=linha, column=col_k).value
        descricao_raw = ws_insumos.cell(row=linha, column=col_d).value
        if tipo_linha == "ítem" or k_raw == "QUANT.":
            grupo_texto = str(descricao_raw).strip() if descricao_raw not in (None, "") else ""
            grupo_atual = grupo_texto or None
            continue

        
        if descricao_raw in (None, "",0,"0") or k_raw in (None, ""):
            continue
        descricao_item = str(descricao_raw).strip()
        if not descricao_item:
            continue
        unidade_raw = ws_insumos.cell(row=linha, column=col_o).value
        if unidade_raw in (None, "",0.0,"0.0",0,"0"):
            continue
        if linha >258:
            quantidade_raw = ws_insumos.cell(row=linha, column=col_j).value    
            if quantidade_raw in (None, "",0.0,"0.0",0,"0"):
                continue
            quantidade_raw = ToDecimal(quantidade_raw, default=Decimal("0.00"))
            l_raw = ws_insumos.cell(row=linha, column=col_l).value
            m_raw = ws_insumos.cell(row=linha, column=col_m).value
            n_raw = ws_insumos.cell(row=linha, column=col_n).value
            multiplicador = 0
            if k_raw == "x" or k_raw == "X":
                multiplicador = 1
            else:
                multiplicador = 0
                continue
            if l_raw == "x" or l_raw == "X":
                multiplicador = 2
            if m_raw == "x" or m_raw == "X":
                multiplicador = 3
            if n_raw == "x" or n_raw == "X":
                multiplicador = 4
            quantidade_raw = quantidade_raw * multiplicador
        else:
            quantidade_raw = ws_insumos.cell(row=linha, column=col_k).value
        valor_raw = ws_insumos.cell(row=linha, column=col_q).value

        if valor_raw in (None, "",0.0,"0.0",0,"0") or quantidade_raw in (None, "",0.0,"0.0",0,"0"):
            continue

        item = ItemOrcamento(
            orcamento_id=orcamento.id,
            descricao_item=descricao_item,
            grupo=grupo_atual,
            material_id=_buscar_material_por_texto_item(descricao_item),
            quantidade=ToDecimal(quantidade_raw, default=Decimal("0.00")),
            valor=ToDecimal(valor_raw, default=Decimal("0.00")),
            unidade=str(unidade_raw).strip() if unidade_raw not in (None, "") else "",
        )
        db.session.add(item)
        itens_importados += 1

    if itens_importados == 0:
        raise ValueError("Nenhum item válido foi encontrado na aba INSUMOS.")

    db.session.commit()
    return {
        "orcamento_id": orcamento.id,
        "nome": orcamento.nome,
        "itens_importados": itens_importados,
        "vinculado_a_existente": orcamento_id_existente is not None,
    }
