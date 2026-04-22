import os
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


def _normalizar_data(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _formatar_duas_casas(valor):
    if valor in (None, ""):
        return ""
    if isinstance(valor, (int, float, Decimal)):
        dec = Decimal(str(valor))
        return str(dec.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    return str(valor).strip()


def ler_template_orcamento(caminho_xlsx):
    wb = load_workbook(filename=caminho_xlsx, data_only=True, read_only=True)

    if "ENTRADAS" not in wb.sheetnames:
        raise ValueError("Aba 'ENTRADAS' não encontrada na planilha.")
    if "INSUMOS" not in wb.sheetnames:
        raise ValueError("Aba 'INSUMOS' não encontrada na planilha.")

    ws_entradas = wb["ENTRADAS"]
    ws_insumos = wb["INSUMOS"]

    nome_orcamento = (ws_entradas["F4"].value or "").strip()
    data_orcamento = _normalizar_data(ws_entradas["O1"].value)

    colunas = ["D", "K", "O", "Q", "R"]
    indices = [column_index_from_string(col) for col in colunas]

    cabecalho = []
    for letra, indice in zip(colunas, indices):
        valor = ws_insumos.cell(row=1, column=indice).value
        cabecalho.append(str(valor).strip() if valor not in (None, "") else letra)

    itens = []
    for linha in range(2, (ws_insumos.max_row or 1) + 1):
        if ws_insumos.cell(row=linha, column=column_index_from_string("D")).value in (None, ""):
            continue

        valores = [ws_insumos.cell(row=linha, column=idx).value for idx in indices]
        if not any(v not in (None, "") for v in valores):
            continue

        descricao_item = str(valores[0]).strip() if valores[0] not in (None, "") else ""
        quantidade = _formatar_duas_casas(valores[2])
        valor = _formatar_duas_casas(valores[4])

        itens.append(
            {
                "descricao_item": descricao_item,
                "quantidade": quantidade,
                "valor": valor,
                "linha_planilha": linha,
                "colunas_origem": dict(zip(colunas, [str(v).strip() if v is not None else "" for v in valores])),
            }
        )

    return {
        "orcamento": {
            "nome": nome_orcamento,
            "data": data_orcamento.isoformat() if data_orcamento else None,
        },
        "cabecalho_insumos": cabecalho,
        "total_itens": len(itens),
        "itens": itens,
    }


def main():
    base_dir = os.path.dirname(__file__)
    caminho = os.path.join(base_dir, "orc.xlsx")
    if not os.path.exists(caminho):
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    resultado = ler_template_orcamento(caminho)
    print(f"Nome: {resultado['orcamento']['nome']}")
    print(f"Data: {resultado['orcamento']['data']}")
    print(f"Itens lidos: {resultado['total_itens']}")
    for item in resultado["itens"][:10]:
        print(f"- Linha {item['linha_planilha']}: {item['descricao_item']} | qtd={item['quantidade']} | valor={item['valor']}")


if __name__ == "__main__":
    main()
